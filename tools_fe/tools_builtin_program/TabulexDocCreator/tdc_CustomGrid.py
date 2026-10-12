# Copyright (C) 2024-present TabulexFlow Team
# SPDX-License-Identifier: AGPL-3.0-or-later

import sys
import os
_tool_dir = os.path.dirname(os.path.abspath(__file__))
if _tool_dir not in sys.path:
    sys.path.insert(0, _tool_dir)

import wx
import wx.grid
import json
import re
import threading
from html import unescape as _html_unescape
from tdc_i18n_manager import translate, translate_list, i18n, get_all_translations
from tdc_logger import log_info, log_error, log_warning, log_debug
from tdc_base_grid import tdc_BaseGrid
from tdc_BaseDocWorker import check_source_file_size
from tfp_toolkit.tfp_base_grid import _UndoEntry
from tfp_toolkit.tfp_dpi import spx_calc

class tdc_CustomGrid(tdc_BaseGrid):
    """TabulexDocCreator 自定义 Grid，继承 tdc_BaseGrid 基类

    基类提供：键盘快捷键（Ctrl+C/V/A/Z/Y, Delete, Enter, Tab）、复制/粘贴、
    撤销/重做（增量 Delta 机制）、行/列标签右键菜单、单元格右键菜单、只读模式等。

    本类保留 BDC 特有的业务逻辑：
    - 4行表头行（HEADER_ROW_SHEET/PLACEHOLDER/DIRECTION/VALUE_COUNT），数据行从 DATA_ROW_START=4 开始
    - 50个数据列，从 COL_DATA_START=4 开始
    - 下拉选择单元格（direction_options, value_count_options）
    - 文件拖放（tdc_FileDropToGrid）
    - 导入/导出功能
    - 自定义只读判断（_is_cell_readonly）
    - 行标签右键菜单（插入行/删除行/追加100行）
    - 单元格右键菜单（导入/导出 + 基类菜单项）
    - 列自动调整宽度
    - 状态栏更新
    - set_fill_method() / set_operation_mode() 动态切换表头可见性
    - get_data_rows() / get_header_data() / get_selected_data_row() 数据提取方法
    - clear_result_column() / set_result() 结果列操作
    """

    COL_ID = 0
    COL_SOURCE_FILE = 1
    COL_SAVE_FILE = 2
    COL_RESULT = 3
    COL_DATA_START = 4
    
    DATA_ROW_START = 4
    
    def __init__(self, parent, tool_instance, doc_type='excel', fill_method='placeholder'):
        super().__init__(parent)
        self.tool_instance = tool_instance
        self.doc_type = doc_type
        self.fill_method = fill_method
        self.operation_mode = 'create_from_template'
        self.is_excel = (doc_type == 'excel')
        if self.is_excel:
            self.HEADER_ROW_SHEET = 0
            self.HEADER_ROW_PLACEHOLDER = 1
            self.HEADER_ROW_DIRECTION = 2
            self.HEADER_ROW_VALUE_COUNT = 3
            self.DATA_ROW_START = 4
            self.header_rows = 4
        else:
            self.HEADER_ROW_TARGET_PAGE = 0
            self.HEADER_ROW_PLACEHOLDER = 1
            self.HEADER_ROW_DIRECTION = 2
            self.HEADER_ROW_VALUE_COUNT = 3
            self.DATA_ROW_START = 4
            self.header_rows = 4
        self.data_columns = 50
        self.read_only_bg_color = wx.Colour(247, 247, 247)
        self.header_bg_color = wx.Colour(230, 230, 230)
        self.header_data_bg_color = wx.Colour(255, 250, 205)
        self.default_col_widths = {}
        self.col_auto_resized = {}
        # 方案6/3（进度可见性）：Worker 线程的网格写回节流 + 当前执行行高亮。
        # 本工具 Worker 原先是在 Worker 线程**直接**调用 grid.SetCellValue / set_result
        # 写网格（未 marshal 到主线程），属跨线程操作 wx 控件：可能界面不刷新、卡顿甚至崩溃。
        # 现改为：非主线程调用 SetCellValue/set_result 时自动进入合并队列，由主线程按事件循环
        # 节拍批量落盘（同名 (row,col) 只保留最后一次值）；并把正在处理的行整行高亮。
        self._pending_lock = threading.Lock()
        self._pending_cell_writes = {}
        self._pending_results = []
        self._pending_highlight_row = -1
        self._flush_scheduled = False
        self._progress_active = False
        self._highlight_row = -1
        self._highlight_saved_colours = {}
        # DPI缩放支持
        self.TFP_DPI_SCALE_FACTOR = getattr(tool_instance, 'TFP_DPI_SCALE_FACTOR', 1.0)
        self._spx = lambda v: spx_calc(v) if v != -1 else -1
        self.default_col_aligns = {
            self.COL_ID: wx.ALIGN_CENTER,
            self.COL_SOURCE_FILE: wx.ALIGN_RIGHT,
            self.COL_SAVE_FILE: wx.ALIGN_RIGHT,
            self.COL_RESULT: wx.ALIGN_LEFT,
        }
        self._init_dropdown_options()
        self._init_grid()
        self._init_menu_ids()
        self._bind_extra_events()
    
    def _init_dropdown_options(self):
        self.direction_options = translate_list('combo_fill_directions')
        self.value_count_options = translate_list('combo_value_counts')
    
    def _init_grid(self):
        total_cols = self.COL_DATA_START + self.data_columns
        self.CreateGrid(self.header_rows + 50, total_cols)
        self._setup_columns()
        self._setup_header_rows()
        self._setup_dropdown_cells()
        self.EnableEditing(True)
        self.DisableDragRowSize()
        self.SetSelectionMode(wx.grid.Grid.GridSelectCells)
        self.SetRowLabelSize(self._spx(100))
        from tdc_FileDrop import tdc_FileDropToGrid
        self.SetDropTarget(tdc_FileDropToGrid(self, self.COL_SOURCE_FILE))
    
    def _setup_columns(self):
        col_labels = {
            self.COL_ID: translate('grid_col_id'),
            self.COL_SOURCE_FILE: translate('grid_col_source_file'),
            self.COL_SAVE_FILE: translate('grid_col_save_file'),
            self.COL_RESULT: translate('grid_col_result'),
        }
        col_widths = {
            self.COL_ID: self._spx(55),
            self.COL_SOURCE_FILE: self._spx(130),
            self.COL_SAVE_FILE: self._spx(130),
            self.COL_RESULT: self._spx(60),
        }
        for col in range(self.COL_DATA_START, self.COL_DATA_START + self.data_columns):
            col_widths[col] = self._spx(100)
            col_labels[col] = f'{translate("grid_col_data")}{col - self.COL_DATA_START + 1}'
        for col, label in col_labels.items():
            self.SetColLabelValue(col, label)
        for col, width in col_widths.items():
            self.SetColSize(col, width)
            self.default_col_widths[col] = width
            self.col_auto_resized[col] = False
    
    def _setup_header_rows(self):
        for row in range(self.header_rows):
            attr = wx.grid.GridCellAttr()
            attr.SetBackgroundColour(self.header_bg_color)
            attr.SetAlignment(wx.ALIGN_CENTER, wx.ALIGN_CENTER)
            self.SetRowAttr(row, attr)
        for row in range(self.header_rows):
            self.SetRowSize(row, self._spx(25))
        self.SetRowSize(self.HEADER_ROW_PLACEHOLDER, self._spx(38))
        for row in range(self.header_rows):
            for col in range(self.COL_RESULT + 1):
                attr = wx.grid.GridCellAttr()
                attr.SetReadOnly(True)
                attr.SetBackgroundColour(self.header_bg_color)
                if col in self.default_col_aligns:
                    attr.SetAlignment(self.default_col_aligns[col], wx.ALIGN_CENTER)
                self.SetAttr(row, col, attr)
        for row in range(self.header_rows):
            for col in range(self.COL_DATA_START, self.COL_DATA_START + self.data_columns):
                attr = wx.grid.GridCellAttr()
                attr.SetBackgroundColour(self.header_data_bg_color)
                attr.SetAlignment(wx.ALIGN_CENTER, wx.ALIGN_CENTER)
                self.SetAttr(row, col, attr)
        self._update_row_labels()
        self._update_header_visibility()
        self._setup_data_cell_alignment()
    
    def _setup_data_cell_alignment(self):
        for row in range(self.DATA_ROW_START, self.GetNumberRows()):
            for col in range(self.COL_RESULT + 1):
                if col in self.default_col_aligns:
                    attr = wx.grid.GridCellAttr()
                    attr.SetAlignment(self.default_col_aligns[col], wx.ALIGN_CENTER)
                    self.SetAttr(row, col, attr)
        for row in range(self.DATA_ROW_START, self.GetNumberRows()):
            for col in range(self.COL_DATA_START, self.COL_DATA_START + self.data_columns):
                attr = wx.grid.GridCellAttr()
                attr.SetAlignment(wx.ALIGN_CENTER, wx.ALIGN_CENTER)
                self.SetAttr(row, col, attr)
    
    def _setup_dropdown_cells(self):
        self._update_dropdown_cells()
    
    def _update_dropdown_cells(self):
        for col in range(self.COL_DATA_START, self.COL_DATA_START + self.data_columns):
            choices = self.direction_options
            attr = wx.grid.GridCellAttr()
            attr.SetEditor(wx.grid.GridCellChoiceEditor(choices))
            attr.SetBackgroundColour(self.header_data_bg_color)
            attr.SetAlignment(wx.ALIGN_CENTER, wx.ALIGN_CENTER)
            self.SetAttr(self.HEADER_ROW_DIRECTION, col, attr)
        for col in range(self.COL_DATA_START, self.COL_DATA_START + self.data_columns):
            attr = wx.grid.GridCellAttr()
            attr.SetEditor(wx.grid.GridCellChoiceEditor(self.value_count_options))
            attr.SetBackgroundColour(self.header_data_bg_color)
            attr.SetAlignment(wx.ALIGN_CENTER, wx.ALIGN_CENTER)
            self.SetAttr(self.HEADER_ROW_VALUE_COUNT, col, attr)
    
    def _update_row_labels(self):
        is_create_mode = self.operation_mode == 'create_from_template'
        if self.is_excel:
            self.SetRowLabelValue(self.HEADER_ROW_SHEET, translate('grid_row_sheet'))
        if not self.is_excel:
            self.SetRowLabelValue(self.HEADER_ROW_TARGET_PAGE, translate('grid_row_target_page'))
        if self.fill_method == 'placeholder':
            placeholder_label = translate('grid_row_placeholder_label')
        else:
            placeholder_label = translate('grid_row_title_label')
        # Create Doc 模式下占位符/标题行为必填，行标签追加“*”
        if is_create_mode:
            placeholder_label += '*'
        self.SetRowLabelValue(self.HEADER_ROW_PLACEHOLDER, placeholder_label)
        self.SetRowLabelValue(self.HEADER_ROW_DIRECTION, translate('grid_row_direction'))
        self.SetRowLabelValue(self.HEADER_ROW_VALUE_COUNT, translate('grid_row_value_count'))
        for row in range(self.DATA_ROW_START, self.GetNumberRows()):
            self.SetRowLabelValue(row, str(row - self.DATA_ROW_START + 1))
    
    def _update_header_visibility(self):
        show_header_rows = True
        is_get_info_mode = self.operation_mode == 'get_file_info'
        if self.doc_type == 'pdf':
            if (is_get_info_mode or
                self.fill_method in ('add_text', 'add_table', 'add_image', 'delete_table')):
                show_header_rows = False
        elif is_get_info_mode:
            show_header_rows = False
        if show_header_rows:
            if self.is_excel:
                self.SetRowSize(self.HEADER_ROW_SHEET, self._spx(25))
            else:
                self.SetRowSize(self.HEADER_ROW_TARGET_PAGE, self._spx(25))
            self.SetRowSize(self.HEADER_ROW_PLACEHOLDER, self._spx(38))
            self.SetRowSize(self.HEADER_ROW_DIRECTION, self._spx(25))
            self.SetRowSize(self.HEADER_ROW_VALUE_COUNT, self._spx(25))
        else:
            if self.is_excel:
                self.SetRowSize(self.HEADER_ROW_SHEET, 0)
            else:
                self.SetRowSize(self.HEADER_ROW_TARGET_PAGE, 0)
            self.SetRowSize(self.HEADER_ROW_PLACEHOLDER, 0)
            self.SetRowSize(self.HEADER_ROW_DIRECTION, 0)
            self.SetRowSize(self.HEADER_ROW_VALUE_COUNT, 0)
        self._update_column_labels_for_pdf_info()
    
    def _update_column_labels_for_pdf_info(self):
        is_get_info_mode = self.operation_mode == 'get_file_info'
        if is_get_info_mode:
            # Get Doc Info 模式：待获取信息文件列同样为必填，标题与 Create 模式一致（Source File*）
            self.SetColLabelValue(self.COL_SOURCE_FILE, translate('grid_col_source_file') + '*')
            
            if self.doc_type == 'pdf':
                info_columns = [
                    'pdf_info_page_count', 'pdf_info_file_size', 'pdf_info_orientation', 
                    'pdf_info_page_size', 'pdf_info_title', 'pdf_info_author', 
                    'pdf_info_creation_date', 'pdf_info_modification_date', 'pdf_info_encrypted'
                ]
            elif self.doc_type == 'excel':
                info_columns = [
                    'excel_info_file_size', 'excel_info_sheet_count', 'excel_info_row_count',
                    'excel_info_col_count', 'excel_info_title', 'excel_info_author',
                    'excel_info_creation_date', 'excel_info_modification_date', 'excel_info_last_modified_by'
                ]
            elif self.doc_type == 'word':
                info_columns = [
                    'word_info_file_size', 'word_info_page_count', 'word_info_word_count',
                    'word_info_char_count', 'word_info_title', 'word_info_author',
                    'word_info_creation_date', 'word_info_modification_date', 'word_info_last_modified_by'
                ]
            elif self.doc_type == 'ppt':
                info_columns = [
                    'ppt_info_file_size', 'ppt_info_slide_count', 'ppt_info_slide_size',
                    'ppt_info_title', 'ppt_info_author', 'ppt_info_creation_date',
                    'ppt_info_modification_date', 'ppt_info_last_modified_by'
                ]
            else:
                info_columns = []
            
            for i, info_key in enumerate(info_columns):
                col = self.COL_DATA_START + i
                if col < self.GetNumberCols():
                    self.SetColLabelValue(col, translate(info_key))
        else:
            # Create Doc 模式下待操作文件与保存文件为必填，列标题追加“*”
            suffix = '*' if self.operation_mode == 'create_from_template' else ''
            self.SetColLabelValue(self.COL_SOURCE_FILE, translate('grid_col_source_file') + suffix)
            self.SetColLabelValue(self.COL_SAVE_FILE, translate('grid_col_save_file') + suffix)
            for i in range(self.data_columns):
                col = self.COL_DATA_START + i
                if col < self.GetNumberCols():
                    self.SetColLabelValue(col, f'{translate("grid_col_data")}{i + 1}')
    
    def _update_col_visibility_for_pdf_info(self):
        is_get_info_mode = self.operation_mode == 'get_file_info'
        if is_get_info_mode:
            self.SetColSize(self.COL_SAVE_FILE, 0)
            self.SetColSize(self.COL_RESULT, 0)
            
            if self.doc_type == 'pdf':
                info_col_count = 9
            elif self.doc_type == 'excel':
                info_col_count = 9
            elif self.doc_type == 'word':
                info_col_count = 9
            elif self.doc_type == 'ppt':
                info_col_count = 8
            else:
                info_col_count = 8
            
            for col in range(self.COL_DATA_START, self.COL_DATA_START + info_col_count):
                if col < self.GetNumberCols():
                    self.SetColSize(col, self._spx(100))
            
            for col in range(self.COL_DATA_START + info_col_count, self.GetNumberCols()):
                self.SetColSize(col, 0)
        else:
            self.SetColSize(self.COL_SAVE_FILE, self._spx(130))
            self.SetColSize(self.COL_RESULT, self._spx(60))
            for col in range(self.COL_DATA_START, self.GetNumberCols()):
                self.SetColSize(col, self._spx(100))
    
    def _apply_info_mode_style(self, is_info):
        """Get Doc Info 模式下，自“File Size”列起的数据列为自动生成的输出列：浅灰且只读"""
        for row in range(self.DATA_ROW_START, self.GetNumberRows()):
            for col in range(self.COL_DATA_START, self.GetNumberCols()):
                attr = wx.grid.GridCellAttr()
                attr.SetAlignment(wx.ALIGN_CENTER, wx.ALIGN_CENTER)
                if is_info:
                    attr.SetReadOnly(True)
                    attr.SetBackgroundColour(self.read_only_bg_color)
                self.SetAttr(row, col, attr)
    
    def set_fill_method(self, fill_method):
        self.fill_method = fill_method
        self._update_row_labels()
        self._update_dropdown_cells()
        self._update_header_visibility()
    
    def set_operation_mode(self, operation_mode):
        mode_changed = (operation_mode != self.operation_mode)
        self.operation_mode = operation_mode
        if mode_changed:
            # 操作模式决定数据列的含义（创建文档=填写定义/数据，获取文档信息=自动生成信息），
            # 切换后旧内容全部作废：清空 ID/Source File/Save File/Result 四列之外所有列的内容
            # （含表头行与数据行），防止旧定义/旧信息被新模式误读
            self._clear_data_columns()
        self._update_row_labels()
        self._update_header_visibility()
        self._update_column_labels_for_pdf_info()
        self._update_col_visibility_for_pdf_info()
        self._setup_data_cell_alignment()
        self._apply_info_mode_style(self.operation_mode == 'get_file_info')

    def count_nonempty_data_cells(self):
        """统计数据列（COL_DATA_START 起，含表头行与数据行）中的非空单元格数量

        用于 Create Doc → Get Doc Info 切换前判断是否弹确认框
        """
        count = 0
        for row in range(self.GetNumberRows()):
            for col in range(self.COL_DATA_START, self.GetNumberCols()):
                if self.GetCellValue(row, col):
                    count += 1
        return count

    def _clear_data_columns(self):
        """清空数据列（COL_DATA_START 起）在全部行中的内容，保留 ID/Source File/Save File/Result 四列"""
        self.BeginBatch()
        try:
            for row in range(self.GetNumberRows()):
                for col in range(self.COL_DATA_START, self.GetNumberCols()):
                    if self.GetCellValue(row, col):
                        self.SetCellValue(row, col, '')
        finally:
            self.EndBatch()
        # 批量清空属破坏性操作，旧撤销记录与新状态不符，清空撤销历史（与导入的处理一致）
        self.clear_undo_history()
    
    def update_fill_method(self, fill_method):
        self.set_fill_method(fill_method)
    
    def _init_menu_ids(self):
        self.ID_IMPORT = wx.NewIdRef()
        self.ID_EXPORT = wx.NewIdRef()
        self.ID_INSERT_ROW = wx.NewIdRef()
        self.ID_DELETE_ROW = wx.NewIdRef()
        self.ID_APPEND_ROWS = wx.NewIdRef()
    
    def _bind_extra_events(self):
        """绑定 BDC 特有的事件（基类已绑定公共事件）"""
        self.Bind(wx.EVT_MENU, self._on_menu_selected)
        self.Bind(wx.grid.EVT_GRID_EDITOR_SHOWN, self._on_editor_shown)
        self.Bind(wx.EVT_SET_FOCUS, self._on_focus)
        self.Bind(wx.grid.EVT_GRID_CELL_LEFT_CLICK, self._on_cell_left_click)
        self.Bind(wx.EVT_KEY_UP, self._on_key_up)
        self.Bind(wx.grid.EVT_GRID_COL_AUTO_SIZE, self._on_col_auto_size)
        # 重绑定 EVT_GRID_SELECT_CELL，覆盖基类的 _on_cell_select，增加状态栏更新
        self.Bind(wx.grid.EVT_GRID_SELECT_CELL, self._on_cell_select_from_base)
        # 基类仅绑定 CHANGING，这里绑定 CHANGED 处理 Get Doc Info 下清空 Source File 的联动
        self.Bind(wx.grid.EVT_GRID_CELL_CHANGED, self._on_cell_changed)
    
    def _update_status_bar(self, row, col):
        if row < 0 or col < 0:
            return
        value = self.GetCellValue(row, col)
        if self.tool_instance and hasattr(self.tool_instance, 'update_status_left'):
            if value:
                display_text = f'{translate("status_current_cell")}: {value}'
            else:
                display_text = f'{translate("status_current_cell")}: '
            self.tool_instance.update_status_left(self.doc_type, display_text)
    
    def _on_key_up(self, event):
        if event.GetKeyCode() in (wx.WXK_UP, wx.WXK_DOWN, wx.WXK_LEFT, wx.WXK_RIGHT, wx.WXK_TAB, wx.WXK_RETURN):
            row = self.GetGridCursorRow()
            col = self.GetGridCursorCol()
            self._update_status_bar(row, col)
        event.Skip()
    
    def _on_focus(self, event):
        row = self.GetGridCursorRow()
        col = self.GetGridCursorCol()
        self._update_status_bar(row, col)
        event.Skip()
    
    def _on_cell_left_click(self, event):
        row = event.GetRow()
        col = event.GetCol()
        self._update_status_bar(row, col)
        event.Skip()
    
    def _on_cell_changed(self, event):
        # Get Doc Info 模式下清空某行 Source File 后，自动清空该行只读信息列，避免残留过期信息
        if self.operation_mode == 'get_file_info' and event.GetCol() == self.COL_SOURCE_FILE:
            row = event.GetRow()
            if not (self.GetCellValue(row, self.COL_SOURCE_FILE) or '').strip():
                self._clear_info_columns_for_rows([row])
        event.Skip()
    
    def _clear_info_columns_for_rows(self, rows):
        """Get Doc Info 模式下清空指定行的只读数据列（程序自动生成的信息列）"""
        if self.operation_mode != 'get_file_info' or not rows:
            return
        self.BeginBatch()
        try:
            for row in sorted(set(rows)):
                if row < self.DATA_ROW_START or row >= self.GetNumberRows():
                    continue
                for col in range(self.COL_DATA_START, self.GetNumberCols()):
                    if self.GetCellValue(row, col):
                        self.SetCellValue(row, col, '')
        finally:
            self.EndBatch()
    
    def _on_cell_select_from_base(self, event):
        """重写基类的 _on_cell_select，增加状态栏更新"""
        # 先调用基类的 hint label 更新
        if self._hint_label is not None:
            col = event.GetCol()
            hint = self._column_hints.get(col, "")
            self._hint_label.SetLabel(hint)
        # 更新状态栏
        row = event.GetRow()
        col = event.GetCol()
        self._update_status_bar(row, col)
        event.Skip()
    
    def _on_editor_shown(self, event):
        row = event.GetRow()
        col = event.GetCol()
        self._update_status_bar(row, col)
        event.Skip()
    
    def _on_col_auto_size(self, event):
        col = event.GetRowOrCol()
        if col == wx.NOT_FOUND:
            event.Skip()
            return
        selected_cols = self._get_selected_columns()
        if selected_cols and len(selected_cols) > 1 and col in selected_cols:
            cols_to_resize = sorted(selected_cols)
        else:
            cols_to_resize = [col]
        reference_col = col
        default_width = self.default_col_widths.get(reference_col, self._spx(100))
        current_width = self.GetColSize(reference_col)
        is_at_default = abs(current_width - default_width) < 2
        if is_at_default:
            for c in cols_to_resize:
                auto_width = self._calculate_auto_width(c)
                self.SetColSize(c, auto_width)
                self.col_auto_resized[c] = True
        else:
            for c in cols_to_resize:
                self.SetColSize(c, self.default_col_widths.get(c, self._spx(100)))
                self.col_auto_resized[c] = False
        event.Veto()
        self.ForceRefresh()
    
    def _get_selected_columns(self):
        selected_cols = set()
        blocks = self.GetSelectionBlockTopLeft()
        if blocks:
            for i, top_left in enumerate(self.GetSelectionBlockTopLeft()):
                bottom_right = self.GetSelectionBlockBottomRight()[i]
                for c in range(top_left[1], bottom_right[1] + 1):
                    selected_cols.add(c)
        return selected_cols
    
    def _calculate_auto_width(self, col):
        max_width = 0
        dc = wx.ClientDC(self)
        font = self.GetDefaultCellFont()
        dc.SetFont(font)
        label = self.GetColLabelValue(col)
        if label:
            label_width = dc.GetTextExtent(label)[0] + self._spx(20)
            max_width = max(max_width, label_width)
        for row in range(self.GetNumberRows()):
            value = self.GetCellValue(row, col)
            if value:
                text_width = dc.GetTextExtent(value)[0] + self._spx(10)
                max_width = max(max_width, text_width)
        return min(max_width, self._spx(500))
    
    # ================================================================
    #  重写基类键盘事件：添加 BDC 特有快捷键
    # ================================================================

    def _on_key_down(self, event):
        """BDC 键盘事件处理：先处理特有快捷键，其余委托基类"""
        # BDC 没有额外的快捷键，全部委托基类处理
        super()._on_key_down(event)

    # ================================================================
    #  重写基类粘贴：保留 BDC 特殊粘贴逻辑
    # ================================================================

    def _read_clipboard_table(self):
        """读取剪贴板表格数据（行×列）。

        优先读取 HTML flavor（text/html）：Excel/LibreOffice 的纯文本 flavor 会**丢弃首尾的全空行**
        （例如选中 A2:H7 且第 2、3 行全空时，纯文本只剩 4 行），导致粘贴后前几行“消失”。
        HTML flavor 会完整保留这些空行，因此优先使用；不可用时回退纯文本。
        """
        if not wx.TheClipboard.Open():
            return None
        try:
            try:
                html_obj = wx.HTMLDataObject()
                if wx.TheClipboard.GetData(html_obj):
                    rows = self._parse_html_table(html_obj.GetHTML())
                    if rows:
                        return rows
            except Exception:
                pass
            text_obj = wx.TextDataObject()
            if wx.TheClipboard.GetData(text_obj):
                return self._parse_tsv_text(text_obj.GetText())
            return None
        finally:
            wx.TheClipboard.Close()

    def _parse_html_table(self, html_text):
        """把剪贴板 HTML 里的 <table> 解析为行×列的纯文本二维列表；无表格时返回 None"""
        if not html_text:
            return None
        table_match = re.search(r'<table\b', html_text, re.IGNORECASE)
        if not table_match:
            return None
        table_html = html_text[table_match.start():]
        rows = []
        for tr_html in re.findall(r'<tr\b[^>]*>(.*?)</tr>', table_html, re.IGNORECASE | re.DOTALL):
            cells = []
            for cell_match in re.finditer(r'<t[dh]\b([^>]*)>(.*?)</t[dh]>',
                                          tr_html, re.IGNORECASE | re.DOTALL):
                attrs, content = cell_match.group(1), cell_match.group(2)
                content = re.sub(r'<br\s*/?>', '\n', content, flags=re.IGNORECASE)
                content = re.sub(r'<[^>]*>', '', content)
                content = _html_unescape(content).replace('\xa0', ' ')
                # 空单元格在 HTML 里常写作 <td><br></td>，此处会得到仅含换行的内容，统一视为空
                if not content.strip():
                    content = ''
                cells.append(content)
                colspan_match = re.search(r'colspan\s*=\s*["\']?(\d+)', attrs, re.IGNORECASE)
                if colspan_match:
                    cells.extend([''] * (int(colspan_match.group(1)) - 1))
            rows.append(cells)
        return rows or None

    def _paste_from_clipboard(self):
        """BDC 粘贴：支持单值多选、可见列粘贴、自动追加行"""
        paste_data = self._read_clipboard_table()
        if not paste_data:
            return
        selected_rows = set()
        selected_cols = set()
        blocks = self.GetSelectionBlockTopLeft()
        if blocks:
            for i, top_left in enumerate(self.GetSelectionBlockTopLeft()):
                bottom_right = self.GetSelectionBlockBottomRight()[i]
                for r in range(top_left[0], bottom_right[0] + 1):
                    selected_rows.add(r)
                for c in range(top_left[1], bottom_right[1] + 1):
                    selected_cols.add(c)
        if not selected_rows or not selected_cols:
            selected_rows.add(self.GetGridCursorRow())
            selected_cols.add(self.GetGridCursorCol())
        single_value = len(paste_data) == 1 and len(paste_data[0]) == 1
        paste_start_row = min(selected_rows)
        paste_start_col = min(selected_cols)
        paste_end_row = paste_start_row
        paste_end_col = paste_start_col

        # 捕获选区用于 Undo
        undo_selection = self._capture_selection()
        undo_changes = []

        if single_value:
            value = paste_data[0][0]
            self.BeginBatch()
            try:
                for r in sorted(selected_rows):
                    for c in sorted(selected_cols):
                        if r < self.GetNumberRows() and c < self.GetNumberCols():
                            if not self._is_cell_readonly(r, c):
                                old_val = self.GetCellValue(r, c)
                                self.SetCellValue(r, c, value)
                                undo_changes.append((r, c, old_val, value))
            finally:
                self.EndBatch()
            paste_end_row = max(selected_rows)
            paste_end_col = max(selected_cols)
        else:
            visible_cols = []
            for col in range(self.GetNumberCols()):
                if self.GetColSize(col) > 0:
                    visible_cols.append(col)

            paste_start_visible_col_idx = 0
            for idx, col in enumerate(visible_cols):
                if col >= paste_start_col:
                    paste_start_visible_col_idx = idx
                    break

            # 与导入一致：仅当粘贴内容首行是本工具“导出”的列标题行时才跳过首行
            # （导出格式含标签列时同时跳过首列），否则不跳过首行，避免普通表格第一行丢失。
            target_cols = visible_cols[paste_start_visible_col_idx:]
            is_header, has_label_col = (
                self._classify_header_row(paste_data[0], target_cols) if paste_data else (False, False)
            )
            if is_header:
                paste_data = paste_data[1:]
                if has_label_col:
                    paste_data = [row[1:] if row else row for row in paste_data]
            elif paste_start_visible_col_idx == 0 and paste_start_row == 0 and self._detect_label_column(paste_data):
                # 从 A2:H7 这类含行标签列的选区复制时，跳过首列（标签列）
                paste_data = [row[1:] if row else row for row in paste_data]
            if not paste_data:
                return

            max_paste_cols = max(len(row) for row in paste_data) if paste_data else 0
            needed_rows = paste_start_row + len(paste_data)
            
            if needed_rows > self.GetNumberRows():
                rows_to_add = needed_rows - self.GetNumberRows()
                self.AppendRows(rows_to_add)
                self._setup_alignment_for_new_rows(self.GetNumberRows() - rows_to_add)
            
            self.BeginBatch()
            try:
                for i, row_data in enumerate(paste_data):
                    target_row = paste_start_row + i
                    if target_row >= self.GetNumberRows():
                        break
                    for j, cell_value in enumerate(row_data):
                        visible_col_idx = paste_start_visible_col_idx + j
                        if visible_col_idx >= len(visible_cols):
                            break
                        target_col = visible_cols[visible_col_idx]
                        if not self._is_cell_readonly(target_row, target_col):
                            old_val = self.GetCellValue(target_row, target_col)
                            self.SetCellValue(target_row, target_col, cell_value)
                            undo_changes.append((target_row, target_col, old_val, cell_value))
            finally:
                self.EndBatch()
            paste_end_row = min(paste_start_row + len(paste_data) - 1, self.GetNumberRows() - 1)
            last_visible_col_idx = min(paste_start_visible_col_idx + max_paste_cols - 1, len(visible_cols) - 1)
            paste_end_col = visible_cols[last_visible_col_idx] if visible_cols and last_visible_col_idx >= 0 else paste_start_col

        self.SelectBlock(paste_start_row, paste_start_col, paste_end_row, paste_end_col)

        if undo_changes:
            entry = _UndoEntry()
            entry.cell_changes = undo_changes
            entry.selection = undo_selection
            self._push_undo(entry)
            pasted_cells = [(r, c) for r, c, _, _ in undo_changes]
            self.on_cells_pasted(pasted_cells)

    # ================================================================
    #  重写基类清空：保留 BDC _is_cell_readonly 判断逻辑
    # ================================================================

    def _clear_selected_cells(self):
        """BDC 清空：使用 _is_cell_readonly 判断只读"""
        if self._read_only:
            return

        selection = self._capture_selection()
        cleared_cells = []
        undo_changes = []

        self.BeginBatch()
        try:
            selected_rows = self.GetSelectedRows()
            selected_cols = self.GetSelectedCols()
            blocks_tl = self.GetSelectionBlockTopLeft()
            blocks_br = self.GetSelectionBlockBottomRight()

            if selected_rows:
                for row in selected_rows:
                    for col in range(self.GetNumberCols()):
                        if not self._is_cell_readonly(row, col):
                            old_val = self.GetCellValue(row, col)
                            self.SetCellValue(row, col, "")
                            cleared_cells.append((row, col))
                            undo_changes.append((row, col, old_val, ""))
            elif selected_cols:
                for col in selected_cols:
                    for row in range(self.GetNumberRows()):
                        if not self._is_cell_readonly(row, col):
                            old_val = self.GetCellValue(row, col)
                            self.SetCellValue(row, col, "")
                            cleared_cells.append((row, col))
                            undo_changes.append((row, col, old_val, ""))
            elif blocks_tl:
                for i in range(len(blocks_tl)):
                    tl = blocks_tl[i]
                    br = blocks_br[i]
                    for row in range(tl[0], br[0] + 1):
                        for col in range(tl[1], br[1] + 1):
                            if not self._is_cell_readonly(row, col):
                                old_val = self.GetCellValue(row, col)
                                self.SetCellValue(row, col, "")
                                cleared_cells.append((row, col))
                                undo_changes.append((row, col, old_val, ""))
            else:
                row = self.GetGridCursorRow()
                col = self.GetGridCursorCol()
                if 0 <= row < self.GetNumberRows() and 0 <= col < self.GetNumberCols():
                    if not self._is_cell_readonly(row, col):
                        old_val = self.GetCellValue(row, col)
                        self.SetCellValue(row, col, "")
                        cleared_cells.append((row, col))
                        undo_changes.append((row, col, old_val, ""))
        finally:
            self.EndBatch()

        if cleared_cells:
            entry = _UndoEntry()
            entry.cell_changes = undo_changes
            entry.selection = selection
            self._push_undo(entry)
            self.on_cells_cleared(cleared_cells)

    # ================================================================
    #  重写基类右键菜单：添加 BDC 特有菜单项
    # ================================================================

    def _on_cell_right_click(self, event):
        """BDC 单元格右键菜单：基类菜单 + BDC 特有项（导入/导出）"""
        if self.IsCellEditControlShown():
            event.Skip()
            return

        self.SetGridCursor(event.GetRow(), event.GetCol())

        menu = wx.Menu()
        # 基类菜单项：复制/粘贴/撤销/重做/删除数据
        self._add_default_cell_menu_items(menu)

        menu.AppendSeparator()

        # BDC 特有：导入、导出
        import_item = menu.Append(self.ID_IMPORT, translate('menu_import'))
        export_item = menu.Append(self.ID_EXPORT, translate('menu_export'))
        self.Bind(wx.EVT_MENU, lambda e: self._import_data(), import_item)
        self.Bind(wx.EVT_MENU, lambda e: self._export_data(), export_item)

        self.PopupMenu(menu)
        menu.Destroy()

    def _on_label_right_click(self, event):
        """BDC 行/列标签右键菜单：BDC 特有的行标签菜单（插入行/删除行/追加100行）"""
        row = event.GetRow()
        col = event.GetCol()
        if row >= 0 and col == -1 and row >= self.DATA_ROW_START:
            menu = wx.Menu()
            menu.Append(self.ID_INSERT_ROW, translate('menu_insert_row'))
            menu.Append(self.ID_DELETE_ROW, translate('menu_delete_row'))
            menu.Append(self.ID_APPEND_ROWS, translate('menu_append_rows'))
            self._context_menu_row = row
            self.PopupMenu(menu)
            menu.Destroy()
        event.Skip()
    
    def _on_menu_selected(self, event):
        menu_id = event.GetId()
        if menu_id == self.ID_IMPORT:
            self._import_data()
        elif menu_id == self.ID_EXPORT:
            self._export_data()
        elif menu_id == self.ID_INSERT_ROW:
            self._insert_row_with_undo()
        elif menu_id == self.ID_DELETE_ROW:
            self._delete_row_with_undo()
        elif menu_id == self.ID_APPEND_ROWS:
            self._append_rows()
        event.Skip()
    
    def _classify_header_row(self, row, visible_cols):
        """判断给定行是否为“列标题行”，返回 (is_header, has_label_col)

        - 导出格式（含标签列）：首格为空，其后单元格均为本工具列标题 → (True, True)
        - 无标签列：首格起即为本工具列标题 → (True, False)
        - 其他（普通数据行）：(False, False)

        用于导入/粘贴时决定是否跳过首行（及首列），避免普通数据表被误删第一行。
        """
        if not row:
            return False, False
        if not str(row[0]).strip() and self._row_all_labels(row, visible_cols, 1):
            return True, True
        if self._row_all_labels(row, visible_cols, 0):
            return True, False
        return False, False

    def _row_all_labels(self, row, visible_cols, offset):
        """从 offset 起逐列比对：非空格子须均为对应列的标题（任意支持语言），且至少匹配到一个"""
        checked = 0
        for index, grid_col in enumerate(visible_cols):
            idx = offset + index
            if idx >= len(row):
                break
            value = str(row[idx]).strip()
            if not value:
                continue
            checked += 1
            if value not in self._column_label_variants(grid_col):
                return False
        return checked > 0

    def _row_label_variants(self, row):
        """Grid 表头行的行标签在全部支持语言下的取值集合（含 '*'）"""
        if row == 0:
            keys = ('grid_row_sheet', 'grid_row_target_page')
        elif row == 1:
            keys = ('grid_row_placeholder_label', 'grid_row_title_label')
        elif row == 2:
            keys = ('grid_row_direction',)
        elif row == 3:
            keys = ('grid_row_value_count',)
        else:
            return set()
        variants = set()
        for key in keys:
            for value in get_all_translations(key):
                variants.add(value)
                variants.add(f'{value}*')
        return variants

    def _detect_label_column(self, rows):
        """判断数据首列是否为 Grid 行标签列（如从 A2:H7 复制/导入时）。

        仅当表头行位置（第 0-3 行）的非空首列值均为对应 Grid 行标签时才成立，
        避免把普通数据表的 ID 列误判为标签列。
        """
        checked = 0
        for index, row in enumerate(rows):
            if index >= self.DATA_ROW_START:
                break
            if not row:
                continue
            value = str(row[0]).strip()
            if not value:
                continue
            checked += 1
            if value not in self._row_label_variants(index):
                return False
        return checked > 0

    def _column_label_variants(self, col):
        """某列标题在所有支持语言下的取值集合（含 Create Doc 模式追加的 '*'）"""
        if col == self.COL_ID:
            keys = ('grid_col_id',)
        elif col == self.COL_SOURCE_FILE:
            keys = ('grid_col_source_file',)
        elif col == self.COL_SAVE_FILE:
            keys = ('grid_col_save_file',)
        elif col == self.COL_RESULT:
            keys = ('grid_col_result',)
        else:
            data_index = col - self.COL_DATA_START + 1
            return {f'{base}{data_index}' for base in get_all_translations('grid_col_data')}
        variants = set()
        for key in keys:
            for value in get_all_translations(key):
                variants.add(value)
                variants.add(f'{value}*')
        return variants

    def _import_data(self):
        dialog = wx.FileDialog(self, translate('menu_import'), wildcard=f'{translate("filetype_excel")} (*.xlsx;*.xls)|*.xlsx;*.xls|{translate("filetype_csv")} (*.csv)|*.csv', style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST)
        if dialog.ShowModal() == wx.ID_OK:
            filepath = dialog.GetPath()
            try:
                if filepath.endswith('.csv'):
                    import csv
                    with open(filepath, 'r', encoding='utf-8-sig') as f:
                        reader = csv.reader(f)
                        rows = list(reader)
                else:
                    import openpyxl
                    # B3-2：解析前限制源文件大小（防资源耗尽）
                    if not check_source_file_size(filepath):
                        log_error(f'Import file too large (>200 MB): {filepath}')
                        dialog.Destroy()
                        wx.MessageBox(translate('msg_file_too_large', size='200 MB'), translate('dialog_title_error'), wx.OK | wx.ICON_ERROR)
                        return
                    wb = openpyxl.load_workbook(filepath, data_only=True)
                    ws = wb.active
                    rows = [[str(cell.value) if cell.value is not None else '' for cell in row] for row in ws.iter_rows()]
                
                if not rows:
                    dialog.Destroy()
                    return
                
                visible_rows = []
                for row in range(self.GetNumberRows()):
                    if self.GetRowSize(row) > 0:
                        visible_rows.append(row)
                
                visible_cols = []
                for col in range(self.GetNumberCols()):
                    if self.GetColSize(col) > 0:
                        visible_cols.append(col)

                # 导入格式：工具“导出”的文件在首行写列标题、首列写行标签。仅当识别到标题行时才跳过
                # 首行/首列；否则按原样从第一行/第一列导入，避免普通数据表的第一行（或第一列 ID）被丢掉。
                is_header, has_label_col = self._classify_header_row(rows[0], visible_cols)
                data_start_row = 1 if is_header else 0
                if is_header:
                    col_offset = 1 if has_label_col else 0
                elif self._detect_label_column(rows):
                    col_offset = 1
                else:
                    col_offset = 0

                if len(rows) > data_start_row:
                    for row_offset, grid_row in enumerate(visible_rows):
                        excel_row = data_start_row + row_offset
                        if excel_row >= len(rows):
                            break
                        excel_row_data = rows[excel_row]
                        for grid_col in visible_cols:
                            excel_col = grid_col + col_offset
                            if excel_col < len(excel_row_data):
                                cell_value = excel_row_data[excel_col]
                                self.SetCellValue(grid_row, grid_col, cell_value)
            except Exception as e:
                wx.MessageBox(f'{translate("msg_error")}: {str(e)}', translate('msg_error'), wx.OK | wx.ICON_ERROR)
            # 导入后清空撤销历史（导入是破坏性操作）
            self.clear_undo_history()
            # Get Doc Info 下导入会替换源文件列表，原有信息列全部作废，统一清空
            if self.operation_mode == 'get_file_info':
                self._clear_info_columns_for_rows(range(self.DATA_ROW_START, self.GetNumberRows()))
        dialog.Destroy()
    
    def _export_data(self):
        tab_labels = {
            'excel': translate('tab_excel'),
            'word': translate('tab_word'),
            'ppt': translate('tab_ppt'),
            'pdf': translate('tab_pdf')
        }
        doc_type = self.doc_type
        tab_label = tab_labels.get(doc_type, doc_type)
        operation_mode = ""
        fill_method = ""
        if self.tool_instance:
            operation_mode = self.tool_instance.get_operation_mode_display(doc_type)
            fill_method = self.tool_instance.get_fill_method_display(doc_type)
        from datetime import datetime
        date_suffix = datetime.now().strftime('_%Y%m%d')
        default_filename = f"{tab_label}_{operation_mode}_{fill_method}{date_suffix}" if operation_mode and fill_method else f"{tab_label}{date_suffix}"
        
        save_dirs = set()
        source_dirs = set()
        for row in range(self.DATA_ROW_START, self.GetNumberRows()):
            if self.GetRowSize(row) > 0:
                save_file = self.GetCellValue(row, self.COL_SAVE_FILE)
                source_file = self.GetCellValue(row, self.COL_SOURCE_FILE)
                if save_file and save_file.strip():
                    save_dir = os.path.dirname(save_file)
                    if save_dir:
                        save_dirs.add(save_dir)
                if source_file and source_file.strip():
                    source_dir = os.path.dirname(source_file)
                    if source_dir:
                        source_dirs.add(source_dir)
        
        default_dir = ""
        if save_dirs:
            if len(save_dirs) == 1:
                default_dir = list(save_dirs)[0]
                parent_dir = os.path.dirname(default_dir)
                if parent_dir and os.path.isdir(parent_dir):
                    default_dir = parent_dir
            else:
                paths = [os.path.abspath(p) for p in save_dirs]
                try:
                    common = os.path.commonpath(paths)
                    if os.path.isdir(common):
                        parent_dir = os.path.dirname(common)
                        if parent_dir and os.path.isdir(parent_dir):
                            default_dir = parent_dir
                        else:
                            default_dir = common
                except Exception:
                    pass
        if not default_dir and source_dirs:
            if len(source_dirs) == 1:
                default_dir = list(source_dirs)[0]
                parent_dir = os.path.dirname(default_dir)
                if parent_dir and os.path.isdir(parent_dir):
                    default_dir = parent_dir
            else:
                paths = [os.path.abspath(p) for p in source_dirs]
                try:
                    common = os.path.commonpath(paths)
                    if os.path.isdir(common):
                        parent_dir = os.path.dirname(common)
                        if parent_dir and os.path.isdir(parent_dir):
                            default_dir = parent_dir
                        else:
                            default_dir = common
                except Exception:
                    pass
        
        if default_filename:
            if default_dir:
                dialog = wx.FileDialog(self, translate('menu_export'), defaultDir=default_dir, defaultFile=default_filename, wildcard=f'{translate("filetype_excel")} (*.xlsx)|*.xlsx|{translate("filetype_csv")} (*.csv)|*.csv', style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT)
            else:
                dialog = wx.FileDialog(self, translate('menu_export'), defaultFile=default_filename, wildcard=f'{translate("filetype_excel")} (*.xlsx)|*.xlsx|{translate("filetype_csv")} (*.csv)|*.csv', style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT)
        else:
            if default_dir:
                dialog = wx.FileDialog(self, translate('menu_export'), defaultDir=default_dir, wildcard=f'{translate("filetype_excel")} (*.xlsx)|*.xlsx|{translate("filetype_csv")} (*.csv)|*.csv', style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT)
            else:
                dialog = wx.FileDialog(self, translate('menu_export'), wildcard=f'{translate("filetype_excel")} (*.xlsx)|*.xlsx|{translate("filetype_csv")} (*.csv)|*.csv', style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT)
        if dialog.ShowModal() == wx.ID_OK:
            filepath = dialog.GetPath()
            try:
                visible_rows = []
                for row in range(self.GetNumberRows()):
                    if self.GetRowSize(row) > 0:
                        visible_rows.append(row)
                
                visible_cols = []
                for col in range(self.GetNumberCols()):
                    if self.GetColSize(col) > 0:
                        visible_cols.append(col)
                
                col_headers = []
                for col in visible_cols:
                    header = self.GetColLabelValue(col)
                    col_headers.append(header)
                
                row_headers = []
                for row in visible_rows:
                    header = self.GetRowLabelValue(row)
                    row_headers.append(header)
                
                if filepath.endswith('.csv'):
                    import csv
                    with open(filepath, 'w', encoding='utf-8-sig', newline='') as f:
                        writer = csv.writer(f)
                        writer.writerow([''] + col_headers)
                        for idx, row in enumerate(visible_rows):
                            row_data = [row_headers[idx]]
                            for col in visible_cols:
                                row_data.append(self.GetCellValue(row, col))
                            writer.writerow(row_data)
                else:
                    from openpyxl import Workbook
                    from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
                    from openpyxl.utils import get_column_letter
                    wb = Workbook()
                    ws = wb.active
                    thin_border = Border(
                        left=Side(style='thin'),
                        right=Side(style='thin'),
                        top=Side(style='thin'),
                        bottom=Side(style='thin')
                    )
                    header_fill = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')
                    header_font = Font(bold=True)
                    
                    header_row = 1
                    ws.cell(row=header_row, column=1, value='')
                    cell = ws.cell(row=header_row, column=1)
                    cell.border = thin_border
                    cell.fill = header_fill
                    cell.font = header_font
                    
                    for col_idx, header in enumerate(col_headers):
                        excel_col = col_idx + 2
                        ws.cell(row=header_row, column=excel_col, value=header)
                        cell = ws.cell(row=header_row, column=excel_col)
                        cell.border = thin_border
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    
                    for row_idx, grid_row in enumerate(visible_rows):
                        excel_row = row_idx + 2
                        ws.cell(row=excel_row, column=1, value=row_headers[row_idx])
                        cell = ws.cell(row=excel_row, column=1)
                        cell.border = thin_border
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        
                        for col_idx, grid_col in enumerate(visible_cols):
                            excel_col = col_idx + 2
                            cell = ws.cell(row=excel_row, column=excel_col)
                            cell.value = self.GetCellValue(grid_row, grid_col)
                            cell.border = thin_border
                            
                            bg_color = self.GetCellBackgroundColour(grid_row, grid_col)
                            if bg_color and bg_color.IsOk():
                                hex_color = bg_color.GetAsString(wx.C2S_HTML_SYNTAX)
                                hex_color = hex_color.lstrip('#')
                                cell.fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
                            
                            horiz, vert = self.GetCellAlignment(grid_row, grid_col)
                            horiz_map = {
                                wx.ALIGN_LEFT: 'left',
                                wx.ALIGN_RIGHT: 'right',
                                wx.ALIGN_CENTER: 'center',
                            }
                            vert_map = {
                                wx.ALIGN_TOP: 'top',
                                wx.ALIGN_BOTTOM: 'bottom',
                                wx.ALIGN_CENTER: 'center',
                            }
                            excel_horiz = horiz_map.get(horiz, 'left')
                            excel_vert = vert_map.get(vert, 'center')
                            cell.alignment = Alignment(horizontal=excel_horiz, vertical=excel_vert)
                    
                    ws.column_dimensions[get_column_letter(1)].width = 10
                    for col_idx, grid_col in enumerate(visible_cols):
                        excel_col = col_idx + 2
                        col_width = self.GetColSize(grid_col)
                        excel_width = max(col_width / 7, 12)
                        ws.column_dimensions[get_column_letter(excel_col)].width = excel_width
                    
                    wb.save(filepath)
            except Exception as e:
                wx.MessageBox(f'{translate("msg_error")}: {str(e)}', translate('msg_error'), wx.OK | wx.ICON_ERROR)
        dialog.Destroy()
    
    def _copy_selection(self):
        """BDC 复制：保留 BFC 的 _format_cell_for_copy 格式化逻辑"""
        selected_rows = set()
        selected_cols = set()
        blocks = self.GetSelectionBlockTopLeft()
        if blocks:
            for i, top_left in enumerate(self.GetSelectionBlockTopLeft()):
                bottom_right = self.GetSelectionBlockBottomRight()[i]
                for r in range(top_left[0], bottom_right[0] + 1):
                    selected_rows.add(r)
                for c in range(top_left[1], bottom_right[1] + 1):
                    selected_cols.add(c)
        if not selected_rows or not selected_cols:
            selected_rows.add(self.GetGridCursorRow())
            selected_cols.add(self.GetGridCursorCol())
        # 跳过隐藏列（宽度为0），避免复制结果里混入隐藏列造成的空列（与导入/粘贴/导出一致）
        visible_cols = [c for c in sorted(selected_cols) if self.GetColSize(c) > 0]
        data = []
        for r in sorted(selected_rows):
            row_data = []
            for c in visible_cols:
                cell_value = self.GetCellValue(r, c)
                row_data.append(self._format_cell_for_copy(cell_value))
            data.append('\t'.join(row_data))
        if wx.TheClipboard.Open():
            wx.TheClipboard.SetData(wx.TextDataObject('\n'.join(data)))
            wx.TheClipboard.Close()
    
    def _format_cell_for_copy(self, value):
        if not value:
            return ''
        needs_quotes = '\n' in value or '\t' in value or '"' in value
        if needs_quotes:
            escaped = value.replace('"', '""')
            return f'"{escaped}"'
        return value
    
    def _is_cell_readonly(self, row, col):
        """BDC 自定义只读判断：表头行+固定列区域为只读；Get Doc Info 模式下的数据列为只读"""
        if row < self.DATA_ROW_START and col < self.COL_DATA_START:
            return True
        if self.operation_mode == 'get_file_info' and col >= self.COL_DATA_START:
            return True
        return False
    
    def _parse_tsv_text(self, text):
        result = []
        current_row = []
        current_cell = ''
        in_quotes = False
        i = 0
        while i < len(text):
            char = text[i]
            if in_quotes:
                if char == '"':
                    if i + 1 < len(text) and text[i + 1] == '"':
                        current_cell += '"'
                        i += 2
                        continue
                    else:
                        in_quotes = False
                        i += 1
                        continue
                else:
                    current_cell += char
                    i += 1
            else:
                if char == '"':
                    in_quotes = True
                    i += 1
                elif char == '\t':
                    current_row.append(current_cell)
                    current_cell = ''
                    i += 1
                elif char == '\r':
                    i += 1
                elif char == '\n':
                    current_row.append(current_cell)
                    result.append(current_row)
                    current_row = []
                    current_cell = ''
                    i += 1
                else:
                    current_cell += char
                    i += 1
        if current_cell or current_row:
            current_row.append(current_cell)
            result.append(current_row)
        return result
    
    def _setup_alignment_for_new_rows(self, start_row):
        for row in range(start_row, self.GetNumberRows()):
            for col in range(self.COL_RESULT + 1):
                if col in self.default_col_aligns:
                    attr = wx.grid.GridCellAttr()
                    attr.SetAlignment(self.default_col_aligns[col], wx.ALIGN_CENTER)
                    self.SetAttr(row, col, attr)
            for col in range(self.COL_DATA_START, self.COL_DATA_START + self.data_columns):
                attr = wx.grid.GridCellAttr()
                attr.SetAlignment(wx.ALIGN_CENTER, wx.ALIGN_CENTER)
                self.SetAttr(row, col, attr)
    
    def _insert_row_with_undo(self):
        """插入行（带撤销支持）"""
        if hasattr(self, '_context_menu_row') and self._context_menu_row is not None:
            row = self._context_menu_row
            self._context_menu_row = None
        else:
            row = self.GetGridCursorRow()
        if row >= self.DATA_ROW_START:
            selection = self._capture_selection()
            self.InsertRows(row, 1)
            self._setup_alignment_for_new_rows(row)
            for r in range(row, self.GetNumberRows()):
                self.SetRowLabelValue(r, str(r - self.DATA_ROW_START + 1))
            # 记录 Undo
            entry = _UndoEntry()
            entry.rows_deleted = [(row, [self.GetCellValue(row, c) for c in range(self.GetNumberCols())])]
            entry.selection = selection
            self._push_undo(entry)
            self.on_rows_added(1)
    
    def _delete_row_with_undo(self):
        """删除行（带撤销支持）"""
        selected_rows = set()
        blocks = self.GetSelectionBlockTopLeft()
        if blocks:
            for i, top_left in enumerate(self.GetSelectionBlockTopLeft()):
                bottom_right = self.GetSelectionBlockBottomRight()[i]
                for r in range(top_left[0], bottom_right[0] + 1):
                    if r >= self.DATA_ROW_START:
                        selected_rows.add(r)
        if not selected_rows:
            if hasattr(self, '_context_menu_row') and self._context_menu_row is not None:
                row = self._context_menu_row
                self._context_menu_row = None
                if row >= self.DATA_ROW_START:
                    selected_rows.add(row)
            else:
                row = self.GetGridCursorRow()
                if row >= self.DATA_ROW_START:
                    selected_rows.add(row)
        if not selected_rows:
            return
        if self.GetNumberRows() - len(selected_rows) < self.DATA_ROW_START + 1:
            return
        
        selection = self._capture_selection()
        # 降序排序，从大到小删除以避免索引偏移
        sorted_rows = sorted(selected_rows, reverse=True)
        # 保存被删除行的完整数据（降序）
        saved_rows = []
        for row in sorted_rows:
            row_data = [self.GetCellValue(row, col) for col in range(self.GetNumberCols())]
            saved_rows.append((row, row_data))
        
        self.BeginBatch()
        try:
            for row in sorted_rows:
                self.DeleteRows(row, 1)
        finally:
            self.EndBatch()
        
        for r in range(self.DATA_ROW_START, self.GetNumberRows()):
            self.SetRowLabelValue(r, str(r - self.DATA_ROW_START + 1))
        
        # 记录 Undo
        entry = _UndoEntry()
        entry.rows_deleted = saved_rows
        entry.selection = selection
        self._push_undo(entry)
        self.on_rows_deleted(sorted_rows)
    
    def _append_rows(self):
        """追加100行（破坏性操作，清空撤销历史）"""
        current_rows = self.GetNumberRows()
        self.AppendRows(100)
        self._setup_alignment_for_new_rows(current_rows)
        for r in range(current_rows, self.GetNumberRows()):
            self.SetRowLabelValue(r, str(r - self.DATA_ROW_START + 1))
        # 追加行后清空撤销历史（涉及行属性设置，难以增量撤销）
        self.clear_undo_history()
    
    # ================================================================
    #  Hook 方法重写
    # ================================================================

    def on_cells_cleared(self, cells):
        """单元格被清空后：更新状态栏；Get Doc Info 下清空 Source File 时同步清空该行信息列"""
        if cells and self.tool_instance and hasattr(self.tool_instance, 'update_status_left'):
            self.tool_instance.update_status_left(self.doc_type, f'{translate("status_current_cell")}: ')
        if self.operation_mode == 'get_file_info' and cells:
            rows = [r for (r, c) in cells if c == self.COL_SOURCE_FILE]
            self._clear_info_columns_for_rows(rows)

    def on_cells_pasted(self, cells):
        """单元格被粘贴后：Get Doc Info 下粘贴 Source File 时同步清空该行信息列"""
        if self.operation_mode == 'get_file_info' and cells:
            rows = [r for (r, c) in cells if c == self.COL_SOURCE_FILE]
            self._clear_info_columns_for_rows(rows)

    def on_rows_added(self, count):
        """行被添加后：设置行标签和对齐"""
        # 行标签和对齐已在 _insert_row_with_undo 中处理
        pass

    def on_rows_deleted(self, rows):
        """行被删除后：更新行标签"""
        # 行标签已在 _delete_row_with_undo 中处理
        pass

    # ================================================================
    #  BDC 特有方法
    # ================================================================

    def get_data_rows(self):
        data = []
        header = self.get_header_data()
        template_usage_count = {}
        is_table_mode = False
        is_delete_table_mode = False
        if self.doc_type == 'pdf':
            if self.fill_method == 'add_table':
                is_table_mode = True
            if self.fill_method == 'delete_table':
                is_delete_table_mode = True
        log_debug(f'get_data_rows: fill_method={self.fill_method}, is_table_mode={is_table_mode}, is_delete_table_mode={is_delete_table_mode}, header_count={len(header)}')
        for row in range(self.DATA_ROW_START, self.GetNumberRows()):
            row_id = self.GetCellValue(row, self.COL_ID)
            source_file = self.GetCellValue(row, self.COL_SOURCE_FILE)
            save_file = self.GetCellValue(row, self.COL_SAVE_FILE)
            # 目标页由每列独立的表头行（item['target_page']）指定，行级 target_page 仅为兼容日志字段
            target_page = ''
            is_inherited = False
            need_lookup = (not (source_file and source_file.strip())) or (not (save_file and save_file.strip()))
            if need_lookup and not is_delete_table_mode:
                for lookup_row in range(row - 1, self.DATA_ROW_START - 1, -1):
                    lookup_source = self.GetCellValue(lookup_row, self.COL_SOURCE_FILE)
                    lookup_save = self.GetCellValue(lookup_row, self.COL_SAVE_FILE)
                    if (lookup_source and lookup_source.strip()) and (lookup_save and lookup_save.strip()):
                        source_file = lookup_source
                        save_file = lookup_save
                        is_inherited = True
                        break
            has_file_data = bool((source_file and source_file.strip()) or (save_file and save_file.strip()))
            has_column_data = False
            for col in header:
                value = self.GetCellValue(row, col)
                if value and value.strip():
                    has_column_data = True
            if is_inherited:
                has_data = has_column_data
            else:
                has_data = has_file_data or (is_table_mode and has_column_data) or (is_delete_table_mode and has_column_data)
            data_items = []
            for col in header:
                value = self.GetCellValue(row, col)
                header_info = header[col]
                item = {
                    'placeholder': header_info.get('placeholder', ''),
                    'direction': header_info.get('direction', ''),
                    'value_count': header_info.get('value_count', ''),
                    'value': value or ''
                }
                if self.is_excel:
                    item['sheet'] = header_info.get('sheet', '')
                else:
                    item['target_page'] = header_info.get('target_page', '')
                data_items.append(item)
            if has_data:
                template_key = f'{source_file}|{save_file}'
                if template_key not in template_usage_count:
                    template_usage_count[template_key] = 0
                template_usage_count[template_key] += 1
                result_value = ''
                if is_inherited:
                    for lookup_row in range(row - 1, self.DATA_ROW_START - 1, -1):
                        lookup_source = self.GetCellValue(lookup_row, self.COL_SOURCE_FILE)
                        lookup_save = self.GetCellValue(lookup_row, self.COL_SAVE_FILE)
                        if (lookup_source and lookup_source.strip()) and (lookup_save and lookup_save.strip()):
                            result_value = self.GetCellValue(lookup_row, self.COL_RESULT) or ''
                            log_debug(f'get_data_rows: row={row}, is_inherited={is_inherited}, lookup_row={lookup_row}, result_value="{result_value}"')
                            break
                else:
                    result_value = self.GetCellValue(row, self.COL_RESULT) or ''
                    log_debug(f'get_data_rows: row={row}, is_inherited={is_inherited}, result_value="{result_value}"')
                row_data = {
                    'row_index': row,
                    'id': row_id or '',
                    'source_file': source_file or '',
                    'save_file': save_file or '',
                    'target_page': target_page or '',
                    'data': data_items,
                    'is_inherited': is_inherited,
                    'multi_value_offset': template_usage_count[template_key],
                    'result': result_value
                }
                data.append(row_data)
        return data
    
    def get_selected_data_row(self):
        row = self.GetGridCursorRow()
        if row >= self.DATA_ROW_START:
            header = self.get_header_data()
            row_id = self.GetCellValue(row, self.COL_ID)
            source_file = self.GetCellValue(row, self.COL_SOURCE_FILE)
            save_file = self.GetCellValue(row, self.COL_SAVE_FILE)
            # 目标页由每列独立的表头行（item['target_page']）指定，行级 target_page 仅为兼容日志字段
            target_page = ''
            is_inherited = False
            is_delete_table_mode = False
            if self.doc_type == 'pdf' and self.fill_method == 'delete_table':
                is_delete_table_mode = True
            need_lookup = (not (source_file and source_file.strip())) or (not (save_file and save_file.strip()))
            if need_lookup and not is_delete_table_mode:
                for lookup_row in range(row - 1, self.DATA_ROW_START - 1, -1):
                    lookup_source = self.GetCellValue(lookup_row, self.COL_SOURCE_FILE)
                    lookup_save = self.GetCellValue(lookup_row, self.COL_SAVE_FILE)
                    if (lookup_source and lookup_source.strip()) and (lookup_save and lookup_save.strip()):
                        source_file = lookup_source
                        save_file = lookup_save
                        is_inherited = True
                        break
            data_items = []
            for col in header:
                value = self.GetCellValue(row, col)
                header_info = header[col]
                item = {
                    'placeholder': header_info.get('placeholder', ''),
                    'direction': header_info.get('direction', ''),
                    'value_count': header_info.get('value_count', ''),
                    'value': value or ''
                }
                if self.is_excel:
                    item['sheet'] = header_info.get('sheet', '')
                else:
                    item['target_page'] = header_info.get('target_page', '')
                data_items.append(item)
            offset = 1
            for lookup_row in range(self.DATA_ROW_START, row):
                lookup_source = self.GetCellValue(lookup_row, self.COL_SOURCE_FILE)
                lookup_save = self.GetCellValue(lookup_row, self.COL_SAVE_FILE)
                if lookup_source == source_file and lookup_save == save_file:
                    offset += 1
                elif not (lookup_source and lookup_source.strip()) and not (lookup_save and lookup_save.strip()):
                    has_lookup_any_data = False
                    for col in header:
                        lookup_val = self.GetCellValue(lookup_row, col)
                        if lookup_val and lookup_val.strip():
                            has_lookup_any_data = True
                            break
                    if has_lookup_any_data:
                        prev_source = None
                        prev_save = None
                        for prev_row in range(lookup_row - 1, self.DATA_ROW_START - 1, -1):
                            prev_source = self.GetCellValue(prev_row, self.COL_SOURCE_FILE)
                            prev_save = self.GetCellValue(prev_row, self.COL_SAVE_FILE)
                            if (prev_source and prev_source.strip()) or (prev_save and prev_save.strip()):
                                break
                        if prev_source == source_file and prev_save == save_file:
                            offset += 1
            return {
                'row_index': row,
                'id': row_id or '',
                'source_file': source_file or '',
                'save_file': save_file or '',
                'target_page': target_page or '',
                'data': data_items,
                'is_inherited': is_inherited,
                'multi_value_offset': offset
            }
        return None
    
    def get_header_data(self):
        header = {}
        need_header_condition = True
        if self.doc_type == 'pdf':
            if self.fill_method in ('add_text', 'add_table', 'add_image', 'delete_table'):
                need_header_condition = False
        for col in range(self.COL_DATA_START, self.GetNumberCols()):
            placeholder = self.GetCellValue(self.HEADER_ROW_PLACEHOLDER, col)
            direction = self.GetCellValue(self.HEADER_ROW_DIRECTION, col)
            value_count = self.GetCellValue(self.HEADER_ROW_VALUE_COUNT, col)
            sheet = ''
            target_page = ''
            if self.is_excel:
                sheet = self.GetCellValue(self.HEADER_ROW_SHEET, col)
            else:
                target_page = self.GetCellValue(self.HEADER_ROW_TARGET_PAGE, col)
            if need_header_condition:
                if placeholder or direction or value_count:
                    header[col] = {
                        'placeholder': placeholder,
                        'direction': direction,
                        'value_count': value_count,
                        'sheet': sheet or '',
                        'target_page': target_page or ''
                    }
            else:
                header[col] = {
                    'placeholder': placeholder,
                    'direction': direction,
                    'value_count': value_count,
                    'sheet': sheet or '',
                    'target_page': target_page or ''
                }
        return header
    
    def get_column_widths(self):
        widths = []
        for col in range(self.COL_DATA_START, self.GetNumberCols()):
            widths.append(self.GetColSize(col))
        return widths
    
    def get_data_row_height(self):
        return self.GetRowSize(self.DATA_ROW_START)
    
    def clear_result_column(self):
        for row in range(self.DATA_ROW_START, self.GetNumberRows()):
            self.SetCellValue(row, self.COL_RESULT, '')
    
    def set_result(self, save_file, result_text):
        """写“执行结果”列；非主线程调用时排队到主线程执行（内部需读取网格定位行）。"""
        if getattr(self, '_pending_lock', None) is not None and not wx.IsMainThread():
            self.queue_result(save_file, result_text)
            return
        self._apply_result(save_file, result_text)

    # ==================== 方案6/3：Worker 写回节流 + 当前执行行高亮 ====================
    # 非主线程调用 SetCellValue/set_result 时自动改为合并队列，由主线程按事件循环节拍
    # 批量落盘（同名 (row,col) 只保留最后一次值）；同时把正在处理的行整行高亮。
    _HIGHLIGHT_COLOUR = wx.Colour(255, 245, 157)

    def begin_operation_progress(self):
        """执行开始时调用（主线程）：进入进度高亮模式。可重复调用。"""
        with self._pending_lock:
            self._progress_active = True
            self._pending_highlight_row = -1
        self._flush_progress()

    def end_operation_progress(self):
        """执行结束时调用（主线程）：同步落盘全部挂起写入并清除当前行高亮。"""
        with self._pending_lock:
            self._progress_active = False
        self._apply_pending()
        self._clear_highlight()

    def SetCellValue(self, row, col, value):
        """非主线程调用时自动改为合并队列写回（线程安全），主线程照常直接写入。

        Worker 线程原先直接写网格属跨线程操作 wx 控件（可能界面不刷新/卡顿/崩溃）；
        此处统一拦截为节流写回，避免事件队列被逐项写入淹没。
        """
        if getattr(self, '_pending_lock', None) is not None and not wx.IsMainThread():
            self.queue_cell_value(row, col, value)
            return
        super().SetCellValue(row, col, value)

    def queue_cell_value(self, row, col, value):
        """Worker 线程调用：单元格写入进入合并队列（线程安全）。"""
        with self._pending_lock:
            self._pending_cell_writes[(row, col)] = value
            self._schedule_flush_locked()

    def queue_result(self, save_file, result_text):
        """Worker 线程调用：执行结果写入进入合并队列（线程安全，主线程解析行号）。"""
        with self._pending_lock:
            self._pending_results.append((save_file, result_text))
            self._schedule_flush_locked()

    def highlight_save_basename(self, basename):
        """主线程调用：按保存文件名（basename）高亮对应数据行（供执行进度逐文件高亮）。"""
        if not basename or not self._progress_active:
            return
        target = basename.strip().lower()
        for row in range(self.DATA_ROW_START, self.GetNumberRows()):
            cell = self.GetCellValue(row, self.COL_SAVE_FILE)
            if cell and os.path.basename(cell.strip()).lower() == target:
                self._set_highlight_row(row)
                return

    def _schedule_flush_locked(self):
        if not self._flush_scheduled:
            self._flush_scheduled = True
            wx.CallAfter(self._flush_progress)

    def _flush_progress(self):
        """主线程：批量落盘挂起写入并更新当前行高亮。"""
        self._apply_pending()
        with self._pending_lock:
            target = self._pending_highlight_row
        if self._progress_active and target >= 0 and target != self._highlight_row:
            self._set_highlight_row(target)

    def _apply_pending(self):
        with self._pending_lock:
            cell_items = list(self._pending_cell_writes.items())
            self._pending_cell_writes.clear()
            results = list(self._pending_results)
            self._pending_results.clear()
            self._flush_scheduled = False
        if not cell_items and not results:
            return
        rows = self.GetNumberRows()
        cols = self.GetNumberCols()
        last_row = -1
        self.BeginBatch()
        try:
            for (row, col), value in cell_items:
                if 0 <= row < rows and 0 <= col < cols:
                    super().SetCellValue(row, col, value)
                if col >= self.COL_DATA_START:
                    last_row = row
            for save_file, text in results:
                r = self._apply_result(save_file, text)
                if r >= 0:
                    last_row = r
        finally:
            self.EndBatch()
        if last_row >= 0:
            with self._pending_lock:
                self._pending_highlight_row = last_row

    def _apply_result(self, save_file, result_text):
        """在主线程把结果写入匹配的数据行；返回写入的行号（-1 表示未匹配）。"""
        row = self._row_for_save_file(save_file)
        if row >= 0:
            super().SetCellValue(row, self.COL_RESULT, result_text)
        return row

    def _row_for_save_file(self, save_file):
        """按保存文件路径定位数据行；与 _group_data_by_save_file 一致（相对路径按源文件目录解析）。"""
        if not save_file:
            return -1
        target = os.path.abspath(save_file.strip())
        for row in range(self.DATA_ROW_START, self.GetNumberRows()):
            cell_save_file = self.GetCellValue(row, self.COL_SAVE_FILE)
            if not cell_save_file or not cell_save_file.strip():
                continue
            resolved = cell_save_file.strip()
            if not os.path.dirname(resolved):
                source_file = self.GetCellValue(row, self.COL_SOURCE_FILE)
                if source_file and source_file.strip():
                    source_dir = os.path.dirname(source_file.strip())
                    if source_dir:
                        resolved = os.path.join(source_dir, resolved)
            if os.path.abspath(resolved) == target:
                return row
        return -1

    def _set_highlight_row(self, row):
        if self._highlight_row == row:
            return
        self._restore_highlight_row()
        if 0 <= row < self.GetNumberRows():
            saved = {}
            for col in range(self.GetNumberCols()):
                saved[col] = self.GetCellBackgroundColour(row, col)
                self.SetCellBackgroundColour(row, col, self._HIGHLIGHT_COLOUR)
            self._highlight_saved_colours = saved
            self._highlight_row = row
            self.RefreshBlock(row, 0, row, self.GetNumberCols() - 1)

    def _restore_highlight_row(self):
        row = self._highlight_row
        if row >= 0 and 0 <= row < self.GetNumberRows() and self._highlight_saved_colours:
            for col, colour in self._highlight_saved_colours.items():
                if col < self.GetNumberCols():
                    self.SetCellBackgroundColour(row, col, colour)
            self.RefreshBlock(row, 0, row, self.GetNumberCols() - 1)
        self._highlight_saved_colours = {}
        self._highlight_row = -1

    def _clear_highlight(self):
        with self._pending_lock:
            self._pending_highlight_row = -1
        self._restore_highlight_row()
