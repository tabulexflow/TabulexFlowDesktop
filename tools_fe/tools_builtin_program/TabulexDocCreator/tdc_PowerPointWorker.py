# Copyright (C) 2024-present TabulexFlow Team
# SPDX-License-Identifier: AGPL-3.0-or-later

import sys
import os
_tool_dir = os.path.dirname(os.path.abspath(__file__))
if _tool_dir not in sys.path:
    sys.path.insert(0, _tool_dir)

import re
import shutil
import tempfile
import wx
from collections import defaultdict
from tdc_i18n_manager import translate
from tdc_logger import log_info, log_error, log_warning, log_debug
from tdc_BaseDocWorker import tdc_BaseDocWorker, check_source_file_size, document_kind_mismatch_message

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

class tdc_PowerPointWorker(tdc_BaseDocWorker):
    worker_type = 'ppt'

    def __init__(self, tool_instance, data, operation_mode, fill_method, params=None):
        super().__init__(tool_instance, data, operation_mode, fill_method, params)
        self._invalid_targets = []
    
    def execute(self):
        log_info(f"[PowerPointWorker] 开始执行, 操作模式: {self.operation_mode}, 填充方式: {self.fill_method}")
        if not HAS_PPTX:
            self._show_error('python-pptx library not installed')
            log_error('[PowerPointWorker] python-pptx library not installed')
            return False
        
        if self.operation_mode == 'get_file_info':
            return self._execute_get_file_info()
        
        grouped_data = self._group_data_by_save_file()
        file_count = len(grouped_data)
        self._total_groups = file_count
        self._skipped_success = 0
        self._skipped_error = 0
        log_info(f"[PowerPointWorker] 待处理文件数: {file_count}")
        success_count = 0
        skipped_count = 0
        for idx, (save_file, rows) in enumerate(grouped_data.items()):
            if self.cancelled:
                break
            self._update_progress(idx + 1, file_count, os.path.basename(save_file))
            should_skip, has_error = self._get_group_result_state(rows)
            if should_skip:
                log_info(f'[PowerPointWorker] Skipping {save_file} - Result column is not empty')
                self._record_skip(has_error)
                skipped_count += 1
                continue
            try:
                self._invalid_targets = []
                result = self._process_file_group(save_file, rows)
                self.log_group_result(save_file, rows, bool(result),
                                      error=None if result else 'process_file_group returned False')
                if result and self._invalid_targets:
                    warning = translate('result_error_invalid_target',
                                        value=', '.join(dict.fromkeys(self._invalid_targets)))
                    self.results.append({'file': save_file, 'success': True, 'warning': warning})
                    success_count += 1
                else:
                    self.results.append({'file': save_file, 'success': result})
                    if result:
                        success_count += 1
            except Exception as e:
                log_error(f'[PowerPointWorker] Error processing {save_file}: {e}')
                self.log_group_result(save_file, rows, False, error=str(e))
                self.results.append({'file': save_file, 'success': False, 'error': str(e)})
        self._clear_progress()
        log_info(f"[PowerPointWorker] 执行完成, 成功: {success_count}/{file_count}, 跳过: {skipped_count}")
        if not self.cancelled:
            self._show_results()
        return True
    
    def _process_file_group(self, save_file, rows):
        if not rows:
            return False
        source_file = rows[0].get('source_file', '')
        if not source_file or not os.path.exists(source_file):
            log_error(f'Source file not found: {source_file}')
            return False
        # B3-2：解析前限制源文件大小（防资源耗尽）
        if not check_source_file_size(source_file):
            log_error(f'Source file too large (>200 MB): {source_file}')
            return False
        # 文件类型与 PowerPoint Tab 不符（如把 Excel/Word 文件填到 PPT Tab）时给出明确提示
        mismatch = document_kind_mismatch_message(source_file, 'ppt')
        if mismatch:
            log_error(f'[PowerPointWorker] Document kind mismatch: {source_file}')
            raise ValueError(mismatch)
        # .ppt（OLE2）不受 python-pptx 支持，与 Word 的 .doc 处理保持一致（规格6.5）
        if self.is_ole2_file(source_file):
            log_error(f'[PowerPointWorker] .ppt format not supported: {source_file}')
            raise ValueError(translate('result_error_ppt_format'))
        save_dir = os.path.dirname(save_file)
        if save_dir and not os.path.exists(save_dir):
            os.makedirs(save_dir, exist_ok=True)
        temp_file = tempfile.NamedTemporaryFile(suffix='.pptx', delete=False)
        temp_path = temp_file.name
        temp_file.close()
        shutil.copy2(source_file, temp_path)
        prs = Presentation(temp_path)
        try:
            self._process_placeholder_mode(prs, rows)
            prs.save(save_file)
            return True
        finally:
            if os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except Exception:
                    pass
    
    def _resolve_target_slides(self, prs, target_page):
        if not target_page or not target_page.strip():
            return list(prs.slides)
        target_page = target_page.strip()
        if target_page.isdigit():
            slide_num = int(target_page)
            if 1 <= slide_num <= len(prs.slides):
                return [prs.slides[slide_num - 1]]
            return []
        # 非法幻灯片编号不再回退为“所有幻灯片”，返回空表示跳过
        return []
    
    def _process_placeholder_mode(self, prs, rows):
        placeholder_values = defaultdict(list)
        placeholder_types = {}
        placeholder_target_pages = {}
        for row in rows:
            is_inherited = row.get('is_inherited', False)
            for item in row.get('data', []):
                placeholder = item.get('placeholder', '')
                value = item.get('value', '')
                value_count = item.get('value_count', '')
                target_page = item.get('target_page', '')
                if placeholder:
                    is_multi_value = 'multi' in value_count.lower() or '多值' in value_count
                    if is_inherited and not is_multi_value:
                        continue
                    placeholder_values[placeholder].append(value)
                    if placeholder not in placeholder_types:
                        placeholder_types[placeholder] = value_count
                    if placeholder not in placeholder_target_pages:
                        placeholder_target_pages[placeholder] = target_page.strip() if target_page else ''
        
        target_pages = set(placeholder_target_pages.values())
        if not target_pages:
            target_pages = {''}
        
        for target_page in target_pages:
            filtered_values = defaultdict(list)
            filtered_types = {}
            for placeholder, values in placeholder_values.items():
                if placeholder_target_pages.get(placeholder, '') == target_page:
                    filtered_values[placeholder] = values
                    if placeholder in placeholder_types:
                        filtered_types[placeholder] = placeholder_types[placeholder]
            
            if not filtered_values:
                continue
            
            placeholder_indices = defaultdict(int)
            target_slides = self._resolve_target_slides(prs, target_page)
            # 规格6.3：非法/超范围幻灯片编号跳过该列，避免误改全部幻灯片
            if target_page and not target_slides:
                log_warning(f'[PowerPointWorker] Invalid or out-of-range target slide "{target_page}", skipped')
                self._invalid_targets.append(str(target_page))
                continue
            
            for slide in target_slides:
                self._process_shapes(slide.shapes, filtered_values, filtered_types, placeholder_indices, slide)
    
    def _process_shapes(self, shapes, placeholder_values, placeholder_types, placeholder_indices, slide):
        """递归处理幻灯片形状：组合形状展开，表格按单元格文本框处理（规格6.2）"""
        for shape in shapes:
            if getattr(shape, 'shapes', None) is not None:
                self._process_shapes(shape.shapes, placeholder_values, placeholder_types, placeholder_indices, slide)
                continue
            if shape.has_text_frame:
                self._process_text_frame(shape.text_frame, placeholder_values, placeholder_types, placeholder_indices, slide, shape)
            if getattr(shape, 'has_table', False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        self._process_text_frame(cell.text_frame, placeholder_values, placeholder_types, placeholder_indices, slide, shape)
    
    def _process_text_frame(self, text_frame, placeholder_values, placeholder_types, placeholder_indices, slide, shape):
        for paragraph in text_frame.paragraphs:
            self._process_paragraph(paragraph, placeholder_values, placeholder_types, placeholder_indices, slide, shape)
    
    def _process_paragraph(self, paragraph, placeholder_values, placeholder_types, placeholder_indices, slide, shape):
        full_text = paragraph.text
        if not full_text:
            return
        pending = {p: v for p, v in placeholder_values.items() if v and p in full_text}
        if not pending:
            return
        
        # 占位符均位于单个 run 内时可局部替换；存在跨 run 占位符时回退为段落级重建
        has_cross_run = any(
            not any(placeholder in run.text for run in paragraph.runs)
            for placeholder in pending
        )
        image_placeholders = set()
        
        if has_cross_run:
            new_text = full_text
            for placeholder, values in pending.items():
                value_count = placeholder_types.get(placeholder, '')
                if 'image' in value_count.lower() or '图片' in value_count:
                    image_placeholders.add(placeholder)
                    new_text = new_text.replace(placeholder, '')
                elif 'multi' in value_count.lower() or '多值' in value_count:
                    # 规格6.2：多值模式下同一段落内的多个占位符依次消费各行值
                    while placeholder in new_text:
                        idx = placeholder_indices[placeholder]
                        if idx >= len(values):
                            break
                        placeholder_indices[placeholder] = idx + 1
                        new_text = new_text.replace(placeholder, str(values[idx]), 1)
                else:
                    new_text = new_text.replace(placeholder, str(values[0]))
            # 段落级重建：把替换后的文本写回首个 run，保留同段其他文字（图片另以浮动形状插入）
            for run in paragraph.runs:
                run.text = ''
            if paragraph.runs:
                paragraph.runs[0].text = new_text
            self._insert_paragraph_images(image_placeholders, pending, slide, shape, paragraph)
            return
        
        # 逐 run 局部替换，避免把整段文本重复写入某个 run
        for run in paragraph.runs:
            if not run.text:
                continue
            for placeholder, values in pending.items():
                if placeholder not in run.text:
                    continue
                value_count = placeholder_types.get(placeholder, '')
                if 'image' in value_count.lower() or '图片' in value_count:
                    image_placeholders.add(placeholder)
                    run.text = run.text.replace(placeholder, '')
                elif 'multi' in value_count.lower() or '多值' in value_count:
                    while placeholder in run.text:
                        idx = placeholder_indices[placeholder]
                        if idx >= len(values):
                            break
                        placeholder_indices[placeholder] = idx + 1
                        run.text = run.text.replace(placeholder, str(values[idx]), 1)
                else:
                    run.text = run.text.replace(placeholder, str(values[0]))
        self._insert_paragraph_images(image_placeholders, pending, slide, shape, paragraph)
    
    def _insert_paragraph_images(self, image_placeholders, pending, slide, shape, paragraph=None):
        for placeholder in image_placeholders:
            values = pending.get(placeholder) or []
            if values and os.path.exists(values[0]):
                self._insert_image_to_slide(slide, shape, values[0], paragraph)
    
    def _insert_image_to_slide(self, slide, shape, img_path, paragraph=None):
        if not HAS_PIL:
            log_error('PIL library not installed for image processing')
            return
        # B3-2：解析前限制图片源文件大小（防资源耗尽）
        if not check_source_file_size(img_path):
            log_error(f'Image file too large (>200 MB): {img_path}')
            return
        try:
            img_params = self.params.get('image_params', {}) or {'width': 200, 'height': 200}
            max_width = img_params.get('width', 200)
            max_height = img_params.get('height', 200)
            pil_img = Image.open(img_path)
            orig_width, orig_height = pil_img.size
            ratio = min(max_width / orig_width, max_height / orig_height)
            if ratio < 1:
                new_width = orig_width * ratio / 96
                new_height = orig_height * ratio / 96
            else:
                new_width = orig_width / 96
                new_height = orig_height / 96
            left = shape.left
            top = self._estimate_paragraph_top(shape, paragraph)
            slide.shapes.add_picture(img_path, left, top, width=Inches(new_width))
        except Exception as e:
            log_error(f'Failed to insert image: {e}')
    
    def _estimate_paragraph_top(self, shape, paragraph):
        """估计占位符所在段落相对形状顶部的纵向偏移，尽量以占位符位置为图片锚点"""
        if paragraph is None or not getattr(shape, 'has_text_frame', False):
            return shape.top
        index = 0
        for i, para in enumerate(shape.text_frame.paragraphs):
            if para._p is paragraph._p:
                index = i
                break
        line_height = None
        for run in paragraph.runs:
            if run.font.size is not None:
                line_height = run.font.size
                break
        if line_height is None:
            line_height = Pt(18)
        # 1.2 为经验行距系数
        return shape.top + int(index * line_height * 1.2)
    
    def _execute_get_file_info(self):
        if not self.data:
            self._show_error(translate('msg_no_data'))
            return False
        
        grid = None
        if self.tool_instance and hasattr(self.tool_instance, 'grids'):
            grid = self.tool_instance.grids.get('ppt')
        
        success_count = 0
        error_count = 0
        
        for idx, row_data in enumerate(self.data):
            source_file = row_data.get('source_file', '')
            if not source_file or not source_file.strip():
                continue
            
            if not os.path.exists(source_file):
                log_warning(f'File not found: {source_file}')
                self.log_info_row_failure(row_data, 'file not found')
                error_count += 1
                continue
            # B3-2：解析前限制源文件大小（防资源耗尽）
            if not check_source_file_size(source_file):
                log_warning(f'Source file too large (>200 MB): {source_file}')
                self.log_info_row_failure(row_data, 'file too large (>200 MB)')
                error_count += 1
                continue
            # 文件类型与 PowerPoint Tab 不符时给出明确提示（规格2.4）
            mismatch = document_kind_mismatch_message(source_file, 'ppt')
            if mismatch:
                log_error(f'[PowerPointWorker] Document kind mismatch, skip get info: {source_file}')
                self.log_info_row_failure(row_data, mismatch)
                error_count += 1
                continue
            # .ppt（OLE2）不受 python-pptx 支持，与执行/预览判定保持一致（规格6.5）
            if self.is_ole2_file(source_file):
                error_msg = translate('result_error_ppt_format')
                log_error(f'[PowerPointWorker] .ppt format not supported, skip get info: {source_file}')
                self.log_info_row_failure(row_data, error_msg)
                error_count += 1
                continue
            try:
                file_size = os.path.getsize(source_file)
                if file_size < 1024:
                    file_size_str = f'{file_size} B'
                elif file_size < 1024 * 1024:
                    file_size_str = f'{file_size / 1024:.1f} KB'
                else:
                    file_size_str = f'{file_size / (1024 * 1024):.1f} MB'
                
                prs = Presentation(source_file)
                slide_count = len(prs.slides)
                
                slide_width = prs.slide_width
                slide_height = prs.slide_height
                slide_width_inches = slide_width / 914400
                slide_height_inches = slide_height / 914400
                slide_size_str = f'{slide_width_inches:.1f}" x {slide_height_inches:.1f}"'
                
                props = prs.core_properties
                title = props.title or ''
                author = props.author or translate('info_not_available')
                created = props.created
                modified = props.modified
                last_modified_by = props.last_modified_by or translate('info_not_available')
                
                if not title:
                    title = os.path.splitext(os.path.basename(source_file))[0]
                
                file_stat = os.stat(source_file)
                file_created = file_stat.st_ctime
                file_modified = file_stat.st_mtime
                
                created_str = translate('pdf_info_not_available')
                if created:
                    try:
                        created_str = created.strftime('%Y-%m-%d')
                    except Exception:
                        pass
                if created_str == translate('pdf_info_not_available'):
                    try:
                        from datetime import datetime
                        created_str = datetime.fromtimestamp(file_created).strftime('%Y-%m-%d')
                    except (ValueError, AttributeError):
                        pass
                
                modified_str = translate('pdf_info_not_available')
                if modified:
                    try:
                        modified_str = modified.strftime('%Y-%m-%d')
                    except Exception:
                        pass
                if modified_str == translate('pdf_info_not_available'):
                    try:
                        from datetime import datetime
                        modified_str = datetime.fromtimestamp(file_modified).strftime('%Y-%m-%d')
                    except (ValueError, AttributeError):
                        pass
                
                info_values = [
                    file_size_str,
                    str(slide_count),
                    slide_size_str,
                    title,
                    author,
                    created_str,
                    modified_str,
                    last_modified_by
                ]
                
                if grid:
                    row_index = row_data.get('row_index', 0)
                    for col_idx, value in enumerate(info_values):
                        grid.SetCellValue(row_index, grid.COL_DATA_START + col_idx, value)
                
                log_info(f'PowerPoint info extracted: {source_file}')
                success_count += 1
                
            except Exception as e:
                log_error(f'Error extracting PowerPoint info from {source_file}: {e}')
                self.log_info_row_failure(row_data, str(e))
                error_count += 1
        
        log_info(f'Get PowerPoint Info completed: {success_count} success, {error_count} errors')
        
        # 获取文档信息：仅在存在未成功的行时弹窗提示，全部成功时静默完成
        if error_count > 0:
            message = translate('msg_info_failed_rows', count=error_count)
            wx.MessageBox(message, translate('msg_warning'), wx.OK | wx.ICON_WARNING)
        
        return True
    
    def _show_results(self):
        import wx
        success_count = sum(1 for r in self.results if r.get('success'))
        total_count = len(self.results)
        if self.tool_instance and hasattr(self.tool_instance, 'grids'):
            grid = self.tool_instance.grids.get('ppt')
            if grid:
                for r in self.results:
                    if r.get('warning'):
                        grid.set_result(r['file'], r['warning'])
                    elif r.get('success'):
                        grid.set_result(r['file'], translate('result_success'))
                    else:
                        error_msg = r.get('error', '')
                        short_error = self._get_short_error(error_msg)
                        # 已知不支持格式等完整提示本身已含“失败”语义，避免重复前缀
                        if error_msg and short_error == error_msg:
                            grid.set_result(r['file'], short_error)
                        else:
                            grid.set_result(r['file'], f'{translate("result_failed")}: {short_error}')
        message = self._build_completion_message(success_count, total_count)
        def _do_show():
            wx.MessageBox(message, translate('tool_name'), wx.OK | wx.ICON_INFORMATION)
        wx.CallAfter(_do_show)
    
    def _get_short_error(self, error_msg):
        if not error_msg:
            return translate('result_error_unknown')
        # .ppt 等已知不支持的格式：直接显示完整提示，避免被归类为泛化错误
        if error_msg == translate('result_error_ppt_format'):
            return error_msg
        # 文件类型与 Tab 不符：直接显示完整提示，避免被归类为泛化错误
        for key in ('msg_wrong_tab_use_word', 'msg_wrong_tab_use_excel', 'msg_wrong_tab_use_ppt'):
            if error_msg == translate(key):
                return error_msg
        error_lower = error_msg.lower()
        if 'permission' in error_lower or '权限' in error_msg:
            return translate('result_error_permission')
        if 'not found' in error_lower or '不存在' in error_msg:
            return translate('result_error_not_found')
        if 'invalid' in error_lower or '无效' in error_msg:
            return translate('result_error_invalid')
        if 'timeout' in error_lower or '超时' in error_msg:
            return translate('result_error_timeout')
        if 'memory' in error_lower or '内存' in error_msg:
            return translate('result_error_memory')
        return translate('result_error_other')
    
    def preview(self, rows_data):
        if not HAS_PPTX:
            return None
        if not rows_data:
            return None
        if isinstance(rows_data, dict):
            rows_data = [rows_data]
        source_file = rows_data[0].get('source_file', '')
        save_file = rows_data[0].get('save_file', '')
        if not source_file or not os.path.exists(source_file):
            self.log_group_result(save_file, rows_data, False, error='Source file not found')
            return None
        # B3-2：解析前限制源文件大小（防资源耗尽）
        if not check_source_file_size(source_file):
            log_error(f'Source file too large (>200 MB), skip preview: {source_file}')
            self.log_group_result(save_file, rows_data, False, error='Source file too large (>200 MB)')
            return None
        # 文件类型与 PowerPoint Tab 不符时给出明确提示（规格2.4）
        mismatch = document_kind_mismatch_message(source_file, 'ppt')
        if mismatch:
            log_error(f'[PowerPointWorker] Document kind mismatch, skip preview: {source_file}')
            self.log_group_result(save_file, rows_data, False, error=mismatch)
            return {'type': 'ppt', 'error': mismatch}
        # .ppt（OLE2）不受 python-pptx 支持，与 Word 的 .doc 处理保持一致（规格6.5）
        if self.is_ole2_file(source_file):
            error_msg = translate('result_error_ppt_format')
            log_error(f'[PowerPointWorker] .ppt format not supported, skip preview: {source_file}')
            self.log_group_result(save_file, rows_data, False, error=error_msg)
            return {'type': 'ppt', 'error': error_msg}
        try:
            prs = Presentation(source_file)
            self._process_placeholder_mode(prs, rows_data)
            slides_data = []
            for slide in prs.slides:
                slide_text = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            slide_text.append(paragraph.text)
                slides_data.append('\n'.join(slide_text))
            self.log_group_result(save_file, rows_data, True)
            return {'type': 'ppt', 'slides': slides_data, 'slide_count': len(prs.slides)}
        except Exception as e:
            log_error(f'Preview failed: {e}')
            self.log_group_result(save_file, rows_data, False, error=str(e))
            return None
    
    def export_slides_as_images(self, row_data, output_dir):
        if not HAS_PPTX or not HAS_PIL:
            return None
        source_file = row_data.get('source_file', '')
        if not source_file or not os.path.exists(source_file):
            return None
        # B3-2：解析前限制源文件大小（防资源耗尽）
        if not check_source_file_size(source_file):
            log_error(f'Source file too large (>200 MB), skip export: {source_file}')
            return None
        try:
            import subprocess
            prs = Presentation(source_file)
            image_paths = []
            return image_paths
        except Exception as e:
            log_error(f'Export slides failed: {e}')
            return None
