# Copyright (C) 2024-present TabulexFlow Team
# SPDX-License-Identifier: AGPL-3.0-or-later

"""TabulexDocCreator Worker 真实逻辑回归测试

覆盖占位符/标题/坐标、合并单元格、多值、目标页、图片、PDF 插入/删除等
历史上出现过的缺陷，避免回归。依赖缺失时自动跳过。
"""
import os
import sys

import pytest

_TOOL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _TOOL_DIR not in sys.path:
    sys.path.insert(0, _TOOL_DIR)


class StubGrid:
    def __init__(self):
        self.COL_DATA_START = 4
        self.info_cells = {}
        self.results = {}

    def SetCellValue(self, row, col, value):
        self.info_cells[(row, col)] = value

    def set_result(self, save_file, text):
        self.results[save_file] = text


class StubTool:
    def __init__(self):
        self.grids = {}

    def get_grid(self, doc_type):
        grid = StubGrid()
        self.grids[doc_type] = grid
        return grid


def make_item(placeholder, value, value_count='Single Value', direction='Down',
              sheet=None, target_page=None):
    item = {'placeholder': placeholder, 'direction': direction,
            'value_count': value_count, 'value': value}
    if sheet is not None:
        item['sheet'] = sheet
    if target_page is not None:
        item['target_page'] = target_page
    return item


def make_row(source, save, items, target_page='', is_inherited=False,
             multi_value_offset=1, row_index=4):
    return {'row_index': row_index, 'id': '1', 'source_file': source,
            'save_file': save, 'target_page': target_page, 'data': items,
            'is_inherited': is_inherited,
            'multi_value_offset': multi_value_offset, 'result': ''}


# ---------------------------------------------------------------- Excel

class TestExcelWorker:
    def _worker(self, tool, rows, fill_method='placeholder', params=None):
        from tdc_ExcelWorker import tdc_ExcelWorker
        return tdc_ExcelWorker(tool, rows, 'create_from_template', fill_method, params or {})

    def test_placeholder_replace_and_formula_skip(self, tmp_path):
        openpyxl = pytest.importorskip('openpyxl')
        src = str(tmp_path / 'src.xlsx')
        dst = str(tmp_path / 'dst.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Sheet1'
        ws['A1'] = 'Hi {{n}}'
        ws['A2'] = '=1+1'
        wb.save(src)
        tool = StubTool()
        tool.get_grid('excel')
        rows = [make_row(src, dst, [make_item('{{n}}', 'Bob', sheet='Sheet1')])]
        assert self._worker(tool, rows).execute()
        out = openpyxl.load_workbook(dst)['Sheet1']
        assert out['A1'].value == 'Hi Bob'
        assert out['A2'].value == '=1+1'

    def test_coordinate_with_sheet_prefix(self, tmp_path):
        openpyxl = pytest.importorskip('openpyxl')
        src = str(tmp_path / 'src.xlsx')
        dst = str(tmp_path / 'dst.xlsx')
        wb = openpyxl.Workbook()
        wb.active.title = 'Sheet1'
        wb.active['A1'] = 'x'
        wb.create_sheet('Sheet2')
        wb.save(src)
        tool = StubTool()
        tool.get_grid('excel')
        # Sheet 表头为 Sheet2，但地址指定 Sheet1!C5，应按地址页名写入 Sheet1
        rows = [make_row(src, dst, [make_item('Sheet1!C5', 'V', direction='Coordinate', sheet='Sheet2')])]
        assert self._worker(tool, rows, 'title_fill').execute()
        wb2 = openpyxl.load_workbook(dst)
        assert wb2['Sheet1']['C5'].value == 'V'
        assert wb2['Sheet2']['C5'].value is None

    def test_coordinate_into_merged_region(self, tmp_path):
        openpyxl = pytest.importorskip('openpyxl')
        src = str(tmp_path / 'src.xlsx')
        dst = str(tmp_path / 'dst.xlsx')
        wb = openpyxl.Workbook()
        wb.active.title = 'Sheet1'
        wb.active.merge_cells('A1:A3')
        wb.save(src)
        tool = StubTool()
        tool.get_grid('excel')
        rows = [make_row(src, dst, [make_item('A3', 'MV', direction='Coordinate', sheet='Sheet1')])]
        worker = self._worker(tool, rows, 'title_fill')
        assert worker.execute()
        assert openpyxl.load_workbook(dst)['Sheet1']['A1'].value == 'MV'

    def test_multi_value_same_cell(self, tmp_path):
        openpyxl = pytest.importorskip('openpyxl')
        src = str(tmp_path / 'src.xlsx')
        dst = str(tmp_path / 'dst.xlsx')
        wb = openpyxl.Workbook()
        wb.active.title = 'Sheet1'
        wb.active['A1'] = '{{v}} {{v}} {{v}}'
        wb.save(src)
        tool = StubTool()
        tool.get_grid('excel')
        rows = [
            make_row(src, dst, [make_item('{{v}}', 'A', value_count='Multi Value', sheet='Sheet1')], multi_value_offset=1),
            make_row('', '', [make_item('{{v}}', 'B', value_count='Multi Value', sheet='Sheet1')], is_inherited=True, multi_value_offset=2),
            make_row('', '', [make_item('{{v}}', 'C', value_count='Multi Value', sheet='Sheet1')], is_inherited=True, multi_value_offset=3),
        ]
        assert self._worker(tool, rows).execute()
        assert openpyxl.load_workbook(dst)['Sheet1']['A1'].value == 'A B C'

    def test_invalid_sheet_is_skipped_with_warning(self, tmp_path):
        openpyxl = pytest.importorskip('openpyxl')
        src = str(tmp_path / 'src.xlsx')
        dst = str(tmp_path / 'dst.xlsx')
        wb = openpyxl.Workbook()
        wb.active.title = 'Sheet1'
        wb.active['A1'] = '{{x}}'
        wb.save(src)
        tool = StubTool()
        tool.get_grid('excel')
        rows = [make_row(src, dst, [make_item('{{x}}', 'V', sheet='Nope')])]
        worker = self._worker(tool, rows)
        worker.execute()
        assert openpyxl.load_workbook(dst)['Sheet1']['A1'].value == '{{x}}'
        assert worker.results[0].get('warning')

    def test_preview_uses_selected_sheet(self, tmp_path):
        openpyxl = pytest.importorskip('openpyxl')
        src = str(tmp_path / 'src.xlsx')
        wb = openpyxl.Workbook()
        wb.active.title = 'Sheet1'
        wb.active['A1'] = 'FIRST'
        ws2 = wb.create_sheet('Sheet2')
        ws2['A1'] = '{{x}}'
        wb.save(src)
        tool = StubTool()
        rows = [make_row(src, str(tmp_path / 'out.xlsx'),
                         [make_item('{{x}}', 'S2VAL', sheet='Sheet2')])]
        worker = self._worker(tool, rows)
        result = worker.preview(rows)
        assert result is not None
        assert result['data'][0][0] == 'S2VAL'

    def test_preview_rejects_word_file_with_clear_message(self, tmp_path):
        pytest.importorskip('docx')
        from docx import Document
        from tdc_i18n_manager import translate
        src = str(tmp_path / 'word.docx')
        Document().save(src)
        tool = StubTool()
        rows = [make_row(src, str(tmp_path / 'out.xlsx'), [make_item('{{x}}', 'V')])]
        worker = self._worker(tool, rows)
        result = worker.preview(rows)
        assert result is not None
        assert result.get('error') == translate('msg_wrong_tab_use_word')

    def test_execute_rejects_word_file_with_clear_message(self, tmp_path):
        pytest.importorskip('docx')
        from docx import Document
        from tdc_i18n_manager import translate
        src = str(tmp_path / 'word.docx')
        Document().save(src)
        tool = StubTool()
        tool.get_grid('excel')
        rows = [make_row(src, str(tmp_path / 'out.xlsx'), [make_item('{{x}}', 'V')])]
        worker = self._worker(tool, rows)
        worker.execute()
        assert worker.results[0]['success'] is False
        assert worker.results[0]['error'] == translate('msg_wrong_tab_use_word')



# ---------------------------------------------------------------- Word

class TestWordWorker:
    def _worker(self, tool, rows, params=None):
        from tdc_WordWorker import tdc_WordWorker
        return tdc_WordWorker(tool, rows, 'create_from_template', 'placeholder', params or {})

    def test_multi_run_and_multi_value(self, tmp_path):
        docx = pytest.importorskip('docx')
        from docx import Document
        src = str(tmp_path / 'src.docx')
        dst = str(tmp_path / 'dst.docx')
        d = Document()
        p = d.add_paragraph()
        p.add_run('Val: ')
        p.add_run('{{v}}')
        p.add_run(' and {{v}}')
        d.save(src)
        tool = StubTool()
        tool.get_grid('word')
        rows = [
            make_row(src, dst, [make_item('{{v}}', 'A', value_count='Multi Value')], multi_value_offset=1),
            make_row('', '', [make_item('{{v}}', 'B', value_count='Multi Value')], is_inherited=True, multi_value_offset=2),
        ]
        self._worker(tool, rows).execute()
        assert Document(dst).paragraphs[0].text == 'Val: A and B'

    def test_cross_run_replace_preserves_formatting(self, tmp_path):
        pytest.importorskip('docx')
        from docx import Document
        from docx.shared import Pt
        src = str(tmp_path / 'src.docx')
        dst = str(tmp_path / 'dst.docx')
        d = Document()
        p = d.add_paragraph()
        r0 = p.add_run('Bold ')
        r0.bold = True
        r0.font.size = Pt(20)
        r1 = p.add_run('{{na')
        r1.italic = True
        p.add_run('me}}')
        r3 = p.add_run(' plain')
        r3.font.size = Pt(10)
        d.save(src)
        tool = StubTool()
        tool.get_grid('word')
        self._worker(tool, [make_row(src, dst, [make_item('{{name}}', 'Bob')])]).execute()
        p2 = Document(dst).paragraphs[0]
        assert p2.text == 'Bold Bob plain'
        assert p2.runs[0].bold is True
        assert p2.runs[0].font.size == Pt(20)
        assert p2.runs[-1].text == ' plain'
        assert p2.runs[-1].font.size == Pt(10)
        assert any(r.text == 'Bob' and r.italic for r in p2.runs)

    def test_image_cross_run_cleared(self, tmp_path):
        pytest.importorskip('docx')
        pytest.importorskip('PIL')
        from docx import Document
        from PIL import Image
        img = str(tmp_path / 'pic.png')
        Image.new('RGB', (60, 40), 'red').save(img)
        src = str(tmp_path / 'src.docx')
        dst = str(tmp_path / 'dst.docx')
        d = Document()
        p = d.add_paragraph()
        p.add_run('A {{pic')
        p.add_run('}} B')
        d.save(src)
        tool = StubTool()
        tool.get_grid('word')
        self._worker(tool, [make_row(src, dst, [make_item('{{pic}}', img, value_count='Image')])],
                     {'image_params': {'width': 200, 'height': 200}}).execute()
        assert '{{pic}}' not in Document(dst).paragraphs[0].text

    def test_target_page_table(self, tmp_path):
        pytest.importorskip('docx')
        from docx import Document
        src = str(tmp_path / 'src.docx')
        dst = str(tmp_path / 'dst.docx')
        d = Document()
        d.add_paragraph('page1')
        d.add_page_break()
        t = d.add_table(rows=1, cols=1)
        t.rows[0].cells[0].paragraphs[0].text = '{{tag}}'
        d.save(src)
        tool = StubTool()
        tool.get_grid('word')
        self._worker(tool, [make_row(src, dst, [make_item('{{tag}}', 'HIT', target_page='2')])]).execute()
        assert Document(dst).tables[0].rows[0].cells[0].text == 'HIT'

    def test_invalid_target_page_skipped(self, tmp_path):
        pytest.importorskip('docx')
        from docx import Document
        src = str(tmp_path / 'src.docx')
        dst = str(tmp_path / 'dst.docx')
        d = Document()
        d.add_paragraph('Keep {{t}}')
        d.save(src)
        tool = StubTool()
        tool.get_grid('word')
        worker = self._worker(tool, [make_row(src, dst, [make_item('{{t}}', 'X', target_page='abc')])])
        worker.execute()
        assert '{{t}}' in Document(dst).paragraphs[0].text
        assert worker.results[0].get('warning')

    def test_preview_ole2_doc_returns_specific_error(self, tmp_path):
        from tdc_WordWorker import tdc_WordWorker
        from tdc_i18n_manager import translate
        src = str(tmp_path / 'legacy.doc')
        with open(src, 'wb') as f:
            f.write(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 64)
        tool = StubTool()
        tool.get_grid('word')
        rows = [make_row(src, str(tmp_path / 'out.doc'), [make_item('{{x}}', 'V')])]
        worker = tdc_WordWorker(tool, rows, 'create_from_template', 'placeholder', {})
        result = worker.preview(rows)
        assert result is not None
        assert result.get('error') == translate('result_error_doc_format')

    def test_get_info_ole2_doc_writes_no_info(self, tmp_path):
        from tdc_WordWorker import tdc_WordWorker
        src = str(tmp_path / 'legacy.doc')
        with open(src, 'wb') as f:
            f.write(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 64)
        tool = StubTool()
        tool.get_grid('word')
        rows = [make_row(src, '', [])]
        worker = tdc_WordWorker(tool, rows, 'get_file_info', 'placeholder', {})
        assert worker.execute() is True
        assert tool.grids['word'].info_cells == {}

    def test_preview_rejects_excel_file_with_clear_message(self, tmp_path):
        openpyxl = pytest.importorskip('openpyxl')
        from tdc_i18n_manager import translate
        src = str(tmp_path / 'sheet.xlsx')
        openpyxl.Workbook().save(src)
        tool = StubTool()
        rows = [make_row(src, str(tmp_path / 'out.docx'), [make_item('{{x}}', 'V')])]
        result = self._worker(tool, rows).preview(rows)
        assert result is not None
        assert result.get('error') == translate('msg_wrong_tab_use_excel')

    def test_execute_rejects_excel_file_with_clear_message(self, tmp_path):
        openpyxl = pytest.importorskip('openpyxl')
        from tdc_i18n_manager import translate
        src = str(tmp_path / 'sheet.xlsx')
        openpyxl.Workbook().save(src)
        tool = StubTool()
        tool.get_grid('word')
        rows = [make_row(src, str(tmp_path / 'out.docx'), [make_item('{{x}}', 'V')])]
        worker = self._worker(tool, rows)
        worker.execute()
        assert worker.results[0]['success'] is False
        assert worker.results[0]['error'] == translate('msg_wrong_tab_use_excel')


# ---------------------------------------------------------------- PowerPoint

class TestPowerPointWorker:
    def _worker(self, tool, rows, params=None):
        from tdc_PowerPointWorker import tdc_PowerPointWorker
        return tdc_PowerPointWorker(tool, rows, 'create_from_template', 'placeholder', params or {})

    def test_multi_run_no_duplication(self, tmp_path):
        pytest.importorskip('pptx')
        from pptx import Presentation
        from pptx.util import Inches
        src = str(tmp_path / 'src.pptx')
        dst = str(tmp_path / 'dst.pptx')
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        p = box.text_frame.paragraphs[0]
        p.add_run().text = 'Hello '
        p.add_run().text = '{{name}}'
        prs.save(src)
        tool = StubTool()
        tool.get_grid('ppt')
        self._worker(tool, [make_row(src, dst, [make_item('{{name}}', 'Bob')])]).execute()
        assert Presentation(dst).slides[0].shapes[0].text_frame.text == 'Hello Bob'

    def test_table_cell_and_multi_value(self, tmp_path):
        pytest.importorskip('pptx')
        from pptx import Presentation
        from pptx.util import Inches
        src = str(tmp_path / 'src.pptx')
        dst = str(tmp_path / 'dst.pptx')
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        gf = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
        gf.table.cell(0, 0).text = '{{a}}'
        gf.table.cell(0, 1).text = '{{a}}'
        prs.save(src)
        tool = StubTool()
        tool.get_grid('ppt')
        rows = [
            make_row(src, dst, [make_item('{{a}}', 'A', value_count='Multi Value')], multi_value_offset=1),
            make_row('', '', [make_item('{{a}}', 'B', value_count='Multi Value')], is_inherited=True, multi_value_offset=2),
        ]
        self._worker(tool, rows).execute()
        table = Presentation(dst).slides[0].shapes[0].table
        assert table.cell(0, 0).text == 'A'
        assert table.cell(0, 1).text == 'B'

    def test_image_removes_placeholder_keeps_text(self, tmp_path):
        pytest.importorskip('pptx')
        pytest.importorskip('PIL')
        from pptx import Presentation
        from pptx.util import Inches
        from PIL import Image
        img = str(tmp_path / 'pic.png')
        Image.new('RGB', (120, 60), 'green').save(img)
        src = str(tmp_path / 'src.pptx')
        dst = str(tmp_path / 'dst.pptx')
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        box.text_frame.paragraphs[0].text = 'keep {{pic}} text'
        prs.save(src)
        tool = StubTool()
        tool.get_grid('ppt')
        self._worker(tool, [make_row(src, dst, [make_item('{{pic}}', img, value_count='Image')])],
                     {'image_params': {'width': 200, 'height': 200}}).execute()
        text = Presentation(dst).slides[0].shapes[0].text_frame.text
        # 图片占位符仅移除占位符本身，同文本框的普通文字保留
        assert '{{pic}}' not in text
        assert 'keep' in text

    def test_image_preserves_other_paragraphs_in_same_textbox(self, tmp_path):
        """图片占位符与文字占位符同处一个文本框时，不得清空其他段落的文字"""
        pytest.importorskip('pptx')
        pytest.importorskip('PIL')
        from pptx import Presentation
        from pptx.util import Inches
        from PIL import Image
        img = str(tmp_path / 'pic.png')
        Image.new('RGB', (120, 60), 'green').save(img)
        src = str(tmp_path / 'src.pptx')
        dst = str(tmp_path / 'dst.pptx')
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        tf = box.text_frame
        tf.paragraphs[0].text = 'This is {{name}} test 01'
        tf.add_paragraph().text = 'This is {{name}} test 02'
        tf.add_paragraph().text = 'Here insert {{pic}}'
        prs.save(src)
        tool = StubTool()
        tool.get_grid('ppt')
        self._worker(tool, [make_row(src, dst, [
            make_item('{{name}}', 'Frame-Top'),
            make_item('{{pic}}', img, value_count='Image'),
        ])], {'image_params': {'width': 200, 'height': 200}}).execute()
        paragraphs = Presentation(dst).slides[0].shapes[0].text_frame.paragraphs
        assert paragraphs[0].text == 'This is Frame-Top test 01'
        assert paragraphs[1].text == 'This is Frame-Top test 02'
        assert '{{pic}}' not in paragraphs[2].text

    def _ole2_ppt(self, tmp_path):
        path = str(tmp_path / 'legacy.ppt')
        with open(path, 'wb') as f:
            f.write(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 64)
        return path

    def test_ppt_ole2_preview_returns_clear_error(self, tmp_path):
        from tdc_i18n_manager import translate
        src = self._ole2_ppt(tmp_path)
        dst = str(tmp_path / 'dst.pptx')
        tool = StubTool()
        tool.get_grid('ppt')
        rows = [make_row(src, dst, [make_item('{{a}}', 'b')])]
        data = self._worker(tool, rows).preview(rows)
        assert data and data.get('error') == translate('result_error_ppt_format')

    def test_ppt_ole2_execute_writes_format_error(self, tmp_path):
        from tdc_i18n_manager import translate
        src = self._ole2_ppt(tmp_path)
        dst = str(tmp_path / 'dst.pptx')
        tool = StubTool()
        grid = tool.get_grid('ppt')
        self._worker(tool, [make_row(src, dst, [make_item('{{a}}', 'b')])]).execute()
        assert grid.results[dst] == translate('result_error_ppt_format')

    def test_ppt_mismatch_reports_other_tab(self, tmp_path):
        openpyxl = pytest.importorskip('openpyxl')
        from tdc_i18n_manager import translate
        src = str(tmp_path / 'book.xlsx')
        openpyxl.Workbook().save(src)
        dst = str(tmp_path / 'dst.pptx')
        tool = StubTool()
        grid = tool.get_grid('ppt')
        rows = [make_row(src, dst, [make_item('{{a}}', 'b')])]
        assert self._worker(tool, rows).preview(rows)['error'] == translate('msg_wrong_tab_use_excel')
        self._worker(tool, rows).execute()
        assert grid.results[dst] == translate('msg_wrong_tab_use_excel')


# ---------------------------------------------------------------- PDF

class TestPDFWorker:
    def _worker(self, tool, rows, fill_method='placeholder', params=None):
        from tdc_PDFWorker import tdc_PDFWorker
        return tdc_PDFWorker(tool, rows, 'create_from_template', fill_method, params or {})

    def _make_pdf(self, path, pages_text):
        fitz = pytest.importorskip('fitz')
        doc = fitz.open()
        for text in pages_text:
            page = doc.new_page()
            page.insert_text((72, 72), text, fontsize=12)
        doc.save(path)
        doc.close()

    def _text(self, path):
        fitz = pytest.importorskip('fitz')
        doc = fitz.open(path)
        text = doc[0].get_text()
        doc.close()
        return text

    def test_placeholder_single_and_long_value(self, tmp_path):
        src = str(tmp_path / 'src.pdf')
        dst = str(tmp_path / 'dst.pdf')
        self._make_pdf(src, ['X {{v}} Y'])
        tool = StubTool()
        long_value = 'A' * 200
        self._worker(tool, [make_row(src, dst, [make_item('{{v}}', long_value)])]).execute()
        assert long_value in self._text(dst).replace('\n', '')

    def test_placeholder_multi_across_pages(self, tmp_path):
        src = str(tmp_path / 'src.pdf')
        dst = str(tmp_path / 'dst.pdf')
        self._make_pdf(src, ['{{v}} p1', '{{v}} p2'])
        tool = StubTool()
        rows = [
            make_row(src, dst, [make_item('{{v}}', 'A', value_count='Multi Value')], multi_value_offset=1),
            make_row('', '', [make_item('{{v}}', 'B', value_count='Multi Value')], is_inherited=True, multi_value_offset=2),
        ]
        self._worker(tool, rows).execute()
        fitz = pytest.importorskip('fitz')
        doc = fitz.open(dst)
        t1, t2 = doc[0].get_text(), doc[1].get_text()
        doc.close()
        assert 'A' in t1 and 'B' in t2

    def test_empty_value_count_defaults_single(self, tmp_path):
        src = str(tmp_path / 'src.pdf')
        dst = str(tmp_path / 'dst.pdf')
        self._make_pdf(src, ['Hi {{a}}'])
        tool = StubTool()
        item = {'placeholder': '{{a}}', 'direction': 'Down', 'value': 'VAL'}
        self._worker(tool, [make_row(src, dst, [item])]).execute()
        assert 'VAL' in self._text(dst)

    def test_insert_text_default(self, tmp_path):
        src = str(tmp_path / 'src.pdf')
        dst = str(tmp_path / 'dst.pdf')
        self._make_pdf(src, ['some content'])
        tool = StubTool()
        rows = [make_row(src, dst, [make_item('', 'Inserted Line 1')])]
        worker = self._worker(tool, rows, 'add_text',
                              {'total_width': 0.015, 'insert_position_index': 3})
        worker.execute()
        assert 'Inserted Line 1' in self._text(dst)

    def test_delete_table_keeps_unrelated_text(self, tmp_path):
        fitz = pytest.importorskip('fitz')
        src = str(tmp_path / 'src.pdf')
        dst = str(tmp_path / 'dst.pdf')
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 60), 'UNRELATED HEADER', fontsize=12)
        x0, y0, w, h = 72, 200, 200, 60
        for i in range(3):
            page.draw_line(fitz.Point(x0, y0 + i * (h / 2)), fitz.Point(x0 + w, y0 + i * (h / 2)))
        for j in range(3):
            page.draw_line(fitz.Point(x0 + j * (w / 2), y0), fitz.Point(x0 + j * (w / 2), y0 + h))
        page.insert_text((x0 + 10, y0 + 25), 'CELLTEXT')
        page.insert_text((72, 700), 'UNRELATED FOOTER', fontsize=12)
        doc.save(src)
        doc.close()
        tool = StubTool()
        rows = [make_row(src, dst, [make_item('', 'CELLTEXT')])]
        self._worker(tool, rows, 'delete_table').execute()
        text = self._text(dst)
        assert 'CELLTEXT' not in text
        assert 'UNRELATED HEADER' in text and 'UNRELATED FOOTER' in text

    def test_encrypted_pdf_is_failure(self, tmp_path):
        fitz = pytest.importorskip('fitz')
        src = str(tmp_path / 'src.pdf')
        dst = str(tmp_path / 'dst.pdf')
        doc = fitz.open()
        doc.new_page().insert_text((72, 72), '{{a}}')
        doc.save(src, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw='u', owner_pw='o')
        doc.close()
        tool = StubTool()
        worker = self._worker(tool, [make_row(src, dst, [make_item('{{a}}', 'X')])])
        worker.execute()
        assert worker.results[0]['success'] is False
        assert worker.results[0].get('warning') is None
