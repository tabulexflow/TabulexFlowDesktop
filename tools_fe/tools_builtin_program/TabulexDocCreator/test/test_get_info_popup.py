# Copyright (C) 2024-present TabulexFlow Team
# SPDX-License-Identifier: AGPL-3.0-or-later

"""Get Doc Info 模式弹窗行为回归测试

需求：全部非空 Source File 行都成功获取信息时不弹窗；
只要存在未成功的行，才弹窗提示未完成的行数。
"""
import os
import sys

import pytest

_TOOL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _TOOL_DIR not in sys.path:
    sys.path.insert(0, _TOOL_DIR)


class StubTool:
    def __init__(self):
        self.grids = {}


def make_row(source):
    return {'row_index': 4, 'id': '1', 'source_file': source, 'save_file': '',
            'target_page': '', 'data': [], 'is_inherited': False,
            'multi_value_offset': 1, 'result': ''}


def _capture_messages(monkeypatch):
    calls = []
    import wx
    monkeypatch.setattr(wx, 'MessageBox', lambda *a, **k: calls.append(a))
    return calls


def _run(worker):
    worker.execute()


def _excel_files(tmp_path):
    openpyxl = pytest.importorskip('openpyxl')
    path = str(tmp_path / 'info.xlsx')
    wb = openpyxl.Workbook()
    wb.active['A1'] = 'x'
    wb.save(path)
    return path


def _word_files(tmp_path):
    pytest.importorskip('docx')
    from docx import Document
    path = str(tmp_path / 'info.docx')
    Document().save(path)
    return path


def _ppt_files(tmp_path):
    pytest.importorskip('pptx')
    from pptx import Presentation
    path = str(tmp_path / 'info.pptx')
    Presentation().save(path)
    return path


def _pdf_files(tmp_path):
    fitz = pytest.importorskip('fitz')
    path = str(tmp_path / 'info.pdf')
    doc = fitz.open()
    doc.new_page()
    doc.save(path)
    doc.close()
    return path


CASES = {
    'excel': ('tdc_ExcelWorker', 'tdc_ExcelWorker', _excel_files, '.xlsx'),
    'word': ('tdc_WordWorker', 'tdc_WordWorker', _word_files, '.docx'),
    'ppt': ('tdc_PowerPointWorker', 'tdc_PowerPointWorker', _ppt_files, '.pptx'),
    'pdf': ('tdc_PDFWorker', 'tdc_PDFWorker', _pdf_files, '.pdf'),
}


@pytest.mark.parametrize('doc_type', list(CASES))
def test_no_popup_when_all_success(monkeypatch, tmp_path, doc_type):
    module_name, class_name, builder, _ext = CASES[doc_type]
    module = __import__(module_name)
    worker_cls = getattr(module, class_name)
    valid = builder(tmp_path)
    calls = _capture_messages(monkeypatch)
    worker = worker_cls(StubTool(), [make_row(valid)], 'get_file_info', 'placeholder', {})
    _run(worker)
    assert calls == [], f'{doc_type}: unexpected popup {calls}'


@pytest.mark.parametrize('doc_type', list(CASES))
def test_popup_when_some_fail(monkeypatch, tmp_path, doc_type):
    module_name, class_name, builder, ext = CASES[doc_type]
    module = __import__(module_name)
    worker_cls = getattr(module, class_name)
    valid = builder(tmp_path)
    missing = str(tmp_path / f'missing{ext}')
    calls = _capture_messages(monkeypatch)
    worker = worker_cls(StubTool(), [make_row(valid), make_row(missing)],
                        'get_file_info', 'placeholder', {})
    _run(worker)
    assert len(calls) == 1, f'{doc_type}: expected 1 popup, got {calls}'
    message = calls[0][0]
    assert '1' in message and 'row' in message.lower()
