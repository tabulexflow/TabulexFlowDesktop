# Copyright (C) 2024-present TabulexFlow Team
# SPDX-License-Identifier: AGPL-3.0-or-later

import sys
import os
_tool_dir = os.path.dirname(os.path.abspath(__file__))
if _tool_dir not in sys.path:
    sys.path.insert(0, _tool_dir)

import re
import shutil
import tempfile
import wx
from collections import defaultdict
from tdc_i18n_manager import translate
from tdc_logger import log_info, log_error, log_warning, log_debug
from tdc_BaseDocWorker import tdc_BaseDocWorker, check_source_file_size, document_kind_mismatch_message

try:
    from docx import Document
    from docx.shared import Inches, Pt
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

class tdc_WordWorker(tdc_BaseDocWorker):
    worker_type = 'word'

    def __init__(self, tool_instance, data, operation_mode, fill_method, params=None):
        super().__init__(tool_instance, data, operation_mode, fill_method, params)
        self._invalid_targets = []
    
    def execute(self):
        log_info(f"[WordWorker] 开始执行, 操作模式: {self.operation_mode}, 填充方式: {self.fill_method}")
        if not HAS_DOCX:
            self._show_error('python-docx library not installed')
            log_error('[WordWorker] python-docx library not installed')
            return False
        
        if self.operation_mode == 'get_file_info':
            return self._execute_get_file_info()
        
        grouped_data = self._group_data_by_save_file()
        file_count = len(grouped_data)
        self._total_groups = file_count
        self._skipped_success = 0
        self._skipped_error = 0
        log_info(f"[WordWorker] 待处理文件数: {file_count}")
        success_count = 0
        skipped_count = 0
        for idx, (save_file, rows) in enumerate(grouped_data.items()):
            if self.cancelled:
                break
            self._update_progress(idx + 1, file_count, os.path.basename(save_file))
            should_skip, has_error = self._get_group_result_state(rows)
            if should_skip:
                log_info(f'[WordWorker] Skipping {save_file} - Result column is not empty')
                self._record_skip(has_error)
                skipped_count += 1
                continue
            try:
                log_debug(f'[WordWorker] Processing file {idx+1}/{file_count}: {save_file}')
                self._invalid_targets = []
                result = self._process_file_group(save_file, rows)
                self.log_group_result(save_file, rows, bool(result),
                                      error=None if result else 'process_file_group returned False')
                if result and self._invalid_targets:
                    warning = translate('result_error_invalid_target',
                                        value=', '.join(dict.fromkeys(self._invalid_targets)))
                    self.results.append({'file': save_file, 'success': True, 'warning': warning})
                    success_count += 1
                else:
                    self.results.append({'file': save_file, 'success': result})
                    if result:
                        success_count += 1
            except Exception as e:
                import traceback
                error_detail = traceback.format_exc()
                log_error(f'[WordWorker] Error processing {save_file}: {e}\n{error_detail}')
                self.log_group_result(save_file, rows, False, error=str(e))
                self.results.append({'file': save_file, 'success': False, 'error': str(e)})
        self._clear_progress()
        log_info(f"[WordWorker] 执行完成, 成功: {success_count}/{file_count}, 跳过: {skipped_count}")
        if not self.cancelled:
            self._show_results()
        return True
    
    def _process_file_group(self, save_file, rows):
        log_debug(f'[WordWorker] _process_file_group start, rows count: {len(rows)}')
        if not rows:
            log_error('[WordWorker] No rows provided')
            return False
        source_file = rows[0].get('source_file', '')
        log_debug(f'[WordWorker] Source file: {source_file}')
        log_debug(f'[WordWorker] Source file exists: {os.path.exists(source_file)}')
        if not source_file or not os.path.exists(source_file):
            log_error(f'Source file not found: {source_file}')
            return False
        # B3-2：解析前限制源文件大小（防资源耗尽）
        if not check_source_file_size(source_file):
            log_error(f'Source file too large (>200 MB): {source_file}')
            return False
        # 文件类型与 Word Tab 不符（如把 Excel/PPT 文件填到 Word Tab）时给出明确提示
        mismatch = document_kind_mismatch_message(source_file, 'word')
        if mismatch:
            log_error(f'[WordWorker] Document kind mismatch: {source_file}')
            raise ValueError(mismatch)
        if self.is_ole2_file(source_file):
            log_error(f'[WordWorker] .doc format not supported: {source_file}')
            raise ValueError(translate('result_error_doc_format'))
        save_dir = os.path.dirname(save_file)
        log_debug(f'[WordWorker] Save directory: {save_dir}')
        if save_dir and not os.path.exists(save_dir):
            log_debug(f'[WordWorker] Creating save directory: {save_dir}')
            os.makedirs(save_dir, exist_ok=True)
        temp_file = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        temp_path = temp_file.name
        temp_file.close()
        log_debug(f'[WordWorker] Temp file path: {temp_path}')
        try:
            log_debug(f'[WordWorker] Copying source file to temp...')
            shutil.copy2(source_file, temp_path)
            log_debug(f'[WordWorker] Copy completed, temp file size: {os.path.getsize(temp_path)} bytes')
        except Exception as copy_err:
            log_error(f'[WordWorker] Failed to copy file: {copy_err}')
            if os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except Exception:
                    pass
            return False
        log_debug(f'[WordWorker] Loading document with python-docx...')
        try:
            doc = Document(temp_path)
            log_debug(f'[WordWorker] Document loaded successfully')
        except Exception as doc_err:
            import traceback
            log_error(f'[WordWorker] Failed to load document: {doc_err}\n{traceback.format_exc()}')
            if os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except Exception:
                    pass
            return False
        try:
            log_debug(f'[WordWorker] Processing placeholder mode...')
            self._process_placeholder_mode(doc, rows)
            log_debug(f'[WordWorker] Saving document to: {save_file}')
            doc.save(save_file)
            log_debug(f'[WordWorker] Document saved successfully')
            return True
        finally:
            if os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except Exception:
                    pass
    
    def _process_placeholder_mode(self, doc, rows):
        log_debug(f'[WordWorker] _process_placeholder_mode start, rows count: {len(rows)}')
        placeholder_values = defaultdict(list)
        placeholder_types = {}
        placeholder_target_pages = {}
        for row in rows:
            is_inherited = row.get('is_inherited', False)
            for item in row.get('data', []):
                placeholder = item.get('placeholder', '')
                value = item.get('value', '')
                value_count = item.get('value_count', '')
                target_page = item.get('target_page', '')
                if placeholder:
                    is_multi_value = 'multi' in value_count.lower() or '多值' in value_count
                    if is_inherited and not is_multi_value:
                        continue
                    placeholder_values[placeholder].append(value)
                    if placeholder not in placeholder_types:
                        placeholder_types[placeholder] = value_count
                    if placeholder not in placeholder_target_pages:
                        placeholder_target_pages[placeholder] = target_page.strip() if target_page else ''
        
        log_debug(f'[WordWorker] Found placeholders: {list(placeholder_values.keys())}')
        # B6-1 更新（2026-09-17）：经确认占位符值即 Grid 中用户输入的数据值，允许完整记录以便复现定位
        log_debug(f'[WordWorker] Placeholder values: {dict(placeholder_values)}')
        
        target_pages = set(placeholder_target_pages.values())
        if not target_pages:
            target_pages = {''}
        
        log_debug(f'[WordWorker] Target pages: {target_pages}')
        
        for target_page in target_pages:
            filtered_values = defaultdict(list)
            filtered_types = {}
            for placeholder, values in placeholder_values.items():
                if placeholder_target_pages.get(placeholder, '') == target_page:
                    filtered_values[placeholder] = values
                    if placeholder in placeholder_types:
                        filtered_types[placeholder] = placeholder_types[placeholder]
            
            if not filtered_values:
                continue
            
            placeholder_indices = defaultdict(int)
            target_paragraphs = self._resolve_target_paragraphs(doc, target_page)
            # 规格5.3：目标页仅支持纯数字；非法/超范围时跳过该列，避免误改整篇文档
            if target_page and not target_paragraphs:
                log_warning(f'[WordWorker] Invalid or out-of-range target page "{target_page}", skipped')
                self._invalid_targets.append(str(target_page))
                continue
            
            log_debug(f'[WordWorker] Processing headers and footers...')
            for section in doc.sections:
                header = section.header
                for paragraph in header.paragraphs:
                    self._process_paragraph(paragraph, filtered_values, filtered_types, placeholder_indices)
                footer = section.footer
                for paragraph in footer.paragraphs:
                    self._process_paragraph(paragraph, filtered_values, filtered_types, placeholder_indices)
            
            log_debug(f'[WordWorker] Processing document body...')
            if target_paragraphs is not None:
                for paragraph in target_paragraphs:
                    self._process_paragraph(paragraph, filtered_values, filtered_types, placeholder_indices)
            else:
                for paragraph in doc.paragraphs:
                    self._process_paragraph(paragraph, filtered_values, filtered_types, placeholder_indices)
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                self._process_paragraph(paragraph, filtered_values, filtered_types, placeholder_indices)
    
    def _get_page_paragraphs(self, doc, page_num):
        paragraphs_by_page = self._split_paragraphs_by_page(doc)
        if 1 <= page_num <= len(paragraphs_by_page):
            return paragraphs_by_page[page_num - 1]
        return []
    
    @staticmethod
    def _iter_block_items(doc):
        """按文档顺序遍历正文段落与表格（表格作为整体块返回）"""
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        for child in doc.element.body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, doc)
            elif isinstance(child, CT_Tbl):
                yield Table(child, doc)
    
    def _split_paragraphs_by_page(self, doc):
        pages = [[]]
        nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        # 规格5.3：按文档顺序包含表格内段落，使指定页时表格内段落也能被筛选
        for block in self._iter_block_items(doc):
            if hasattr(block, 'rows'):
                paragraphs = [p for row in block.rows for cell in row.cells for p in cell.paragraphs]
            else:
                paragraphs = [block]
            for paragraph in paragraphs:
                pages[-1].append(paragraph)
                has_page_break = False
                for run in paragraph.runs:
                    br_elements = run._element.findall('.//w:br', nsmap)
                    for br in br_elements:
                        br_type = br.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type', '')
                        if br_type == 'page':
                            has_page_break = True
                            break
                    if has_page_break:
                        break
                if has_page_break:
                    pages.append([])
        last_section = None
        for section in doc.sections:
            if last_section is not None:
                section_start = section.start_type
                if section_start is not None and section_start != 2:
                    pages.append([])
            last_section = section
        return [p for p in pages if p]
    
    def _resolve_target_paragraphs(self, doc, target_page):
        if not target_page:
            return None
        
        if target_page.isdigit():
            page_num = int(target_page)
            page_paragraphs = self._get_page_paragraphs(doc, page_num)
            if page_paragraphs:
                return page_paragraphs
            return []
        
        # 非数字目标页不受支持（规格5.3），返回空列表表示跳过而非处理所有页
        return []
    
    def _process_paragraph(self, paragraph, placeholder_values, placeholder_types, placeholder_indices):
        for placeholder, values in placeholder_values.items():
            if not values:
                continue
            value_count = placeholder_types.get(placeholder, '')
            if 'image' in value_count.lower() or '图片' in value_count:
                if placeholder in paragraph.text:
                    img_path = values[0]
                    if os.path.exists(img_path):
                        self._clear_placeholder_in_runs(paragraph, placeholder)
                        self._insert_image_to_paragraph(paragraph, img_path)
            else:
                if 'multi' in value_count.lower() or '多值' in value_count:
                    # 规格5.2：多值模式下同一段落内的多个占位符依次消费各行值
                    while placeholder in paragraph.text:
                        idx = placeholder_indices[placeholder]
                        if idx >= len(values):
                            break
                        new_value = values[idx]
                        placeholder_indices[placeholder] = idx + 1
                        self._replace_in_runs(paragraph, placeholder, str(new_value), replace_all=False)
                else:
                    if placeholder in paragraph.text:
                        new_value = values[0]
                        self._replace_in_runs(paragraph, placeholder, str(new_value), replace_all=True)

    def _replace_in_runs(self, paragraph, placeholder, new_value, replace_all=True):
        if not paragraph.runs:
            return
        if replace_all:
            # 逐个替换（含跨 run），保留占位符之外文本的原有格式
            while placeholder in paragraph.text:
                before = paragraph.text
                if not self._replace_cross_run_once(paragraph, placeholder, new_value):
                    break
                if paragraph.text == before:
                    break
        else:
            self._replace_cross_run_once(paragraph, placeholder, new_value)
    
    def _replace_cross_run_once(self, paragraph, placeholder, new_value):
        """替换文档顺序中的第一个占位符

        占位符跨 run 时仅合并占位符实际覆盖的 run（新值继承起始 run 的格式），
        其余 run 的文本与格式保持不变。
        """
        runs = paragraph.runs
        if not runs:
            return False
        texts = [run.text for run in runs]
        full_text = ''.join(texts)
        pos = full_text.find(placeholder)
        if pos < 0:
            return False
        end = pos + len(placeholder)
        cumulative = 0
        start_run = end_run = None
        start_offset = end_offset = 0
        for idx, text in enumerate(texts):
            next_cumulative = cumulative + len(text)
            if start_run is None and pos < next_cumulative:
                start_run = idx
                start_offset = pos - cumulative
            if end <= next_cumulative:
                end_run = idx
                end_offset = end - cumulative
                break
            cumulative = next_cumulative
        if start_run is None or end_run is None:
            return False
        if start_run == end_run:
            run = runs[start_run]
            run.text = texts[start_run][:start_offset] + new_value + texts[start_run][end_offset:]
            return True
        prefix = texts[start_run][:start_offset]
        suffix = texts[end_run][end_offset:]
        runs[start_run].text = prefix + new_value
        for idx in range(start_run + 1, end_run):
            runs[idx].text = ''
        runs[end_run].text = suffix
        return True

    def _clear_placeholder_in_runs(self, paragraph, placeholder):
        self._replace_in_runs(paragraph, placeholder, '', replace_all=True)

    def _insert_image_to_paragraph(self, paragraph, img_path):
        if not HAS_PIL:
            log_error('PIL library not installed for image processing')
            return
        # B3-2：解析前限制图片源文件大小（防资源耗尽）
        if not check_source_file_size(img_path):
            log_error(f'Image file too large (>200 MB): {img_path}')
            return
        try:
            img_params = self.params.get('image_params', {}) or {'width': 200, 'height': 200}
            max_width = img_params.get('width', 200)
            max_height = img_params.get('height', 200)
            pil_img = Image.open(img_path)
            orig_width, orig_height = pil_img.size
            ratio = min(max_width / orig_width, max_height / orig_height)
            if ratio < 1:
                new_width = orig_width * ratio / 96
                new_height = orig_height * ratio / 96
            else:
                new_width = orig_width / 96
                new_height = orig_height / 96
            run = paragraph.add_run()
            run.add_picture(img_path, width=Inches(new_width))
        except Exception as e:
            log_error(f'Failed to insert image: {e}')
    
    def _execute_get_file_info(self):
        if not self.data:
            self._show_error(translate('msg_no_data'))
            return False
        
        grid = None
        if self.tool_instance and hasattr(self.tool_instance, 'grids'):
            grid = self.tool_instance.grids.get('word')
        
        success_count = 0
        error_count = 0
        
        for idx, row_data in enumerate(self.data):
            source_file = row_data.get('source_file', '')
            if not source_file or not source_file.strip():
                continue
            
            if not os.path.exists(source_file):
                log_warning(f'File not found: {source_file}')
                self.log_info_row_failure(row_data, 'file not found')
                error_count += 1
                continue
            # B3-2：解析前限制源文件大小（防资源耗尽）
            if not check_source_file_size(source_file):
                log_warning(f'Source file too large (>200 MB): {source_file}')
                self.log_info_row_failure(row_data, 'file too large (>200 MB)')
                error_count += 1
                continue
            # 文件类型与 Word Tab 不符时给出明确提示
            mismatch = document_kind_mismatch_message(source_file, 'word')
            if mismatch:
                log_error(f'[WordWorker] Document kind mismatch, skip get info: {source_file}')
                self.log_info_row_failure(row_data, mismatch)
                error_count += 1
                continue
            # .doc（OLE2）不受 python-docx 支持，与 execute() 的判定保持一致（规格5.x）
            if self.is_ole2_file(source_file):
                error_msg = translate('result_error_doc_format')
                log_error(f'[WordWorker] .doc format not supported, skip get info: {source_file}')
                self.log_info_row_failure(row_data, error_msg)
                error_count += 1
                continue
            try:
                file_size = os.path.getsize(source_file)
                if file_size < 1024:
                    file_size_str = f'{file_size} B'
                elif file_size < 1024 * 1024:
                    file_size_str = f'{file_size / 1024:.1f} KB'
                else:
                    file_size_str = f'{file_size / (1024 * 1024):.1f} MB'
                
                doc = Document(source_file)
                
                page_count = 0
                word_count = 0
                char_count = 0
                
                for para in doc.paragraphs:
                    text = para.text
                    word_count += len(text.split())
                    char_count += len(text)
                
                page_count = translate('word_page_count_unavailable')
                
                props = doc.core_properties
                title = props.title or ''
                author = props.author or translate('info_not_available')
                created = props.created
                modified = props.modified
                last_modified_by = props.last_modified_by or translate('info_not_available')
                
                if not title:
                    title = os.path.splitext(os.path.basename(source_file))[0]
                
                file_stat = os.stat(source_file)
                file_created = file_stat.st_ctime
                file_modified = file_stat.st_mtime
                
                created_str = translate('pdf_info_not_available')
                if created:
                    try:
                        created_str = created.strftime('%Y-%m-%d')
                    except Exception:
                        pass
                if created_str == translate('pdf_info_not_available'):
                    try:
                        from datetime import datetime
                        created_str = datetime.fromtimestamp(file_created).strftime('%Y-%m-%d')
                    except (ValueError, AttributeError):
                        pass
                
                modified_str = translate('pdf_info_not_available')
                if modified:
                    try:
                        modified_str = modified.strftime('%Y-%m-%d')
                    except Exception:
                        pass
                if modified_str == translate('pdf_info_not_available'):
                    try:
                        from datetime import datetime
                        modified_str = datetime.fromtimestamp(file_modified).strftime('%Y-%m-%d')
                    except (ValueError, AttributeError):
                        pass
                
                info_values = [
                    file_size_str,
                    page_count,
                    str(word_count),
                    str(char_count),
                    title,
                    author,
                    created_str,
                    modified_str,
                    last_modified_by
                ]
                
                if grid:
                    row_index = row_data.get('row_index', 0)
                    for col_idx, value in enumerate(info_values):
                        grid.SetCellValue(row_index, grid.COL_DATA_START + col_idx, value)
                
                log_info(f'Word info extracted: {source_file}')
                success_count += 1
                
            except Exception as e:
                log_error(f'Error extracting Word info from {source_file}: {e}')
                self.log_info_row_failure(row_data, str(e))
                error_count += 1
        
        log_info(f'Get Word Info completed: {success_count} success, {error_count} errors')
        
        # 获取文档信息：仅在存在未成功的行时弹窗提示，全部成功时静默完成
        if error_count > 0:
            message = translate('msg_info_failed_rows', count=error_count)
            wx.MessageBox(message, translate('msg_warning'), wx.OK | wx.ICON_WARNING)
        
        return True
    
    def _show_results(self):
        success_count = sum(1 for r in self.results if r.get('success'))
        total_count = len(self.results)
        if self.tool_instance and hasattr(self.tool_instance, 'grids'):
            grid = self.tool_instance.grids.get('word')
            if grid:
                for r in self.results:
                    if r.get('warning'):
                        grid.set_result(r['file'], r['warning'])
                    elif r.get('success'):
                        grid.set_result(r['file'], translate('result_success'))
                    else:
                        error_msg = r.get('error', '')
                        short_error = self._get_short_error(error_msg)
                        grid.set_result(r['file'], f'{translate("result_failed")}: {short_error}')
        message = self._build_completion_message(success_count, total_count)
        def _do_show():
            wx.MessageBox(message, translate('tool_name'), wx.OK | wx.ICON_INFORMATION)
        wx.CallAfter(_do_show)
    
    def _get_short_error(self, error_msg):
        if not error_msg:
            return translate('result_error_unknown')
        # 文件类型与 Tab 不符：直接显示完整提示，避免被归类为泛化错误
        for key in ('msg_wrong_tab_use_word', 'msg_wrong_tab_use_excel', 'msg_wrong_tab_use_ppt'):
            if error_msg == translate(key):
                return error_msg
        error_lower = error_msg.lower()
        if 'permission' in error_lower or '权限' in error_msg:
            return translate('result_error_permission')
        if 'not found' in error_lower or '不存在' in error_msg:
            return translate('result_error_not_found')
        if 'invalid' in error_lower or '无效' in error_msg:
            return translate('result_error_invalid')
        if 'timeout' in error_lower or '超时' in error_msg:
            return translate('result_error_timeout')
        if 'memory' in error_lower or '内存' in error_msg:
            return translate('result_error_memory')
        return translate('result_error_other')
    
    def preview(self, rows_data):
        if not HAS_DOCX:
            return None
        if not rows_data:
            return None
        if isinstance(rows_data, dict):
            rows_data = [rows_data]
        source_file = rows_data[0].get('source_file', '')
        save_file = rows_data[0].get('save_file', '')
        if not source_file or not os.path.exists(source_file):
            self.log_group_result(save_file, rows_data, False, error='Source file not found')
            return None
        # B3-2：解析前限制源文件大小（防资源耗尽）
        if not check_source_file_size(source_file):
            log_error(f'Source file too large (>200 MB), skip preview: {source_file}')
            self.log_group_result(save_file, rows_data, False, error='Source file too large (>200 MB)')
            return None
        # 文件类型与 Word Tab 不符（如把 Excel/PPT 文件填到 Word Tab）时给出明确提示
        mismatch = document_kind_mismatch_message(source_file, 'word')
        if mismatch:
            log_error(f'[WordWorker] Document kind mismatch, skip preview: {source_file}')
            self.log_group_result(save_file, rows_data, False, error=mismatch)
            return {'type': 'word', 'error': mismatch}
        # .doc（OLE2）不受 python-docx 支持，与 execute() 的判定保持一致（规格5.x）
        if self.is_ole2_file(source_file):
            error_msg = translate('result_error_doc_format')
            log_error(f'[WordWorker] .doc format not supported, skip preview: {source_file}')
            self.log_group_result(save_file, rows_data, False, error=error_msg)
            return {'type': 'word', 'error': error_msg}
        try:
            doc = Document(source_file)
            self._process_placeholder_mode(doc, rows_data)
            text_content = []
            for paragraph in doc.paragraphs:
                text_content.append(paragraph.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            text_content.append(paragraph.text)
            for section in doc.sections:
                header = section.header
                for paragraph in header.paragraphs:
                    text_content.append(paragraph.text)
                footer = section.footer
                for paragraph in footer.paragraphs:
                    text_content.append(paragraph.text)
            self.log_group_result(save_file, rows_data, True)
            return {'type': 'word', 'text': '\n'.join(text_content)}
        except Exception as e:
            log_error(f'Preview failed: {e}')
            self.log_group_result(save_file, rows_data, False, error=str(e))
            return None
