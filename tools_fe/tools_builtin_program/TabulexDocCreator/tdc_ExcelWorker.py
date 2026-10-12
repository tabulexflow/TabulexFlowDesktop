# Copyright (C) 2024-present TabulexFlow Team
# SPDX-License-Identifier: AGPL-3.0-or-later

import sys
import os
_tool_dir = os.path.dirname(os.path.abspath(__file__))
if _tool_dir not in sys.path:
    sys.path.insert(0, _tool_dir)

import re
import shutil
import tempfile
import zipfile
import wx
from collections import defaultdict
from copy import copy
from xml.etree import ElementTree as ET
from tdc_i18n_manager import translate
from tdc_logger import log_info, log_error, log_warning, log_debug
from tdc_BaseDocWorker import tdc_BaseDocWorker, check_source_file_size, document_kind_mismatch_message

try:
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False

class tdc_ExcelWorker(tdc_BaseDocWorker):
    worker_type = 'excel'

    def __init__(self, tool_instance, data, operation_mode, fill_method, params=None):
        super().__init__(tool_instance, data, operation_mode, fill_method, params)
        self._invalid_targets = []
    
    def execute(self):
        log_info(f"[ExcelWorker] 开始执行, 操作模式: {self.operation_mode}, 填充方式: {self.fill_method}")
        if not HAS_OPENPYXL:
            self._show_error('openpyxl library not installed')
            log_error('[ExcelWorker] openpyxl library not installed')
            return False
        
        if self.operation_mode == 'get_file_info':
            return self._execute_get_file_info()
        
        grouped_data = self._group_data_by_save_file()
        file_count = len(grouped_data)
        self._total_groups = file_count
        self._skipped_success = 0
        self._skipped_error = 0
        log_info(f"[ExcelWorker] 待处理文件数: {file_count}")
        success_count = 0
        skipped_count = 0
        for idx, (save_file, rows) in enumerate(grouped_data.items()):
            if self.cancelled:
                break
            self._update_progress(idx + 1, file_count, os.path.basename(save_file))
            should_skip, has_error = self._get_group_result_state(rows)
            if should_skip:
                log_info(f'[ExcelWorker] Skipping {save_file} - Result column is not empty')
                self._record_skip(has_error)
                skipped_count += 1
                continue
            try:
                self._invalid_targets = []
                result = self._process_file_group(save_file, rows)
                self.log_group_result(save_file, rows, bool(result),
                                      error=None if result else 'process_file_group returned False')
                if result and self._invalid_targets:
                    warning = translate('result_error_invalid_target',
                                        value=', '.join(dict.fromkeys(self._invalid_targets)))
                    self.results.append({'file': save_file, 'success': True, 'warning': warning})
                    success_count += 1
                else:
                    self.results.append({'file': save_file, 'success': result})
                    if result:
                        success_count += 1
            except Exception as e:
                log_error(f'[ExcelWorker] Error processing {save_file}: {e}')
                self.log_group_result(save_file, rows, False, error=str(e))
                self.results.append({'file': save_file, 'success': False, 'error': str(e)})
        self._clear_progress()
        log_info(f"[ExcelWorker] 执行完成, 成功: {success_count}/{file_count}, 跳过: {skipped_count}")
        if not self.cancelled:
            self._show_results()
        return True
    
    def _check_macros(self, filepath):
        """检测Excel文件是否包含VBA宏（仅记录警告，不阻止处理，因为openpyxl本身不执行宏）"""
        try:
            from oletools.olevba import detect_vba_macros
            if detect_vba_macros(filepath):
                log_warning(f'文件包含VBA宏: {filepath}')
                return True
        except Exception:
            pass  # oletools不可用或文件格式不支持，跳过检测
        return False
    
    @staticmethod
    def _truncate_xml_after_root(data):
        """截取 XML 首个根元素（含）之前的内容，丢弃根元素之后的多余内容。

        用于修复部分模板中畸形文档属性 XML（根元素结束后残留重复片段）。
        无法定位根元素时返回 None。
        """
        text = data.decode('utf-8', errors='replace')
        i, n = 0, len(text)
        while i < n:
            if text[i] != '<':
                i += 1
                continue
            if text.startswith('<?', i) or text.startswith('<!', i):
                end = text.find('>', i)
                if end == -1:
                    return None
                i = end + 1
                continue
            break
        match = re.match(r'<([A-Za-z_][\w:.\-]*)', text[i:])
        if not match:
            return None
        tag = match.group(1)
        open_pattern = re.compile(r'<' + re.escape(tag) + r'(?=[\s/>])')
        close_tag = '</' + tag
        depth = 0
        pos = i
        while pos < n:
            if text.startswith(close_tag, pos):
                depth -= 1
                end = text.find('>', pos)
                if end == -1:
                    return None
                pos = end + 1
                if depth == 0:
                    return text[:pos].encode('utf-8')
            elif open_pattern.match(text, pos):
                end = text.find('>', pos)
                if end == -1:
                    return None
                if text[end - 1] == '/':
                    pos = end + 1
                else:
                    depth += 1
                    pos = end + 1
            else:
                pos += 1
        return None
    
    def _repair_docprops_xml(self, path):
        """修复 xlsx 中畸形的 docProps/*.xml（如根元素结束后残留多余内容）。

        某些模板的 docProps/core.xml 在根元素之后残留重复片段，会使 openpyxl
        解析失败（lxml XMLSyntaxError）。这里仅重写解析失败的 docProps XML，
        其余 zip 内容原样保留。返回是否执行了修复。
        """
        try:
            with zipfile.ZipFile(path, 'r') as zin:
                names = set(zin.namelist())
                targets = [n for n in ('docProps/core.xml', 'docProps/app.xml', 'docProps/custom.xml')
                           if n in names]
                if not targets:
                    return False
                fixed = {}
                for name in targets:
                    raw = zin.read(name)
                    try:
                        ET.fromstring(raw)
                    except ET.ParseError:
                        repaired = self._truncate_xml_after_root(raw)
                        if repaired is None:
                            continue
                        try:
                            ET.fromstring(repaired)
                        except ET.ParseError:
                            continue
                        fixed[name] = repaired
                if not fixed:
                    return False
                tmp_path = path + '.repair'
                with zipfile.ZipFile(tmp_path, 'w', zipfile.ZIP_DEFLATED) as zout:
                    for item in zin.infolist():
                        data = fixed.get(item.filename)
                        if data is None:
                            data = zin.read(item.filename)
                        zout.writestr(item, data)
            os.replace(tmp_path, path)
            log_warning(f'Repaired malformed docProps XML in working copy: {path} ({sorted(fixed)})')
            return True
        except Exception as e:
            log_error(f'Failed to repair docProps XML in {path}: {e}')
            return False
    
    def _make_repaired_copy(self, source_file):
        """复制源文件到临时文件并修复畸形 docProps XML，返回临时文件路径"""
        temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        temp_path = temp_file.name
        temp_file.close()
        shutil.copy2(source_file, temp_path)
        self._repair_docprops_xml(temp_path)
        return temp_path
    
    def _load_workbook_repairing(self, source_file, **kwargs):
        """加载工作簿；若因 docProps XML 畸形失败，则复制并修复后重试。

        返回 (workbook, temp_path)；temp_path 非空时由调用方在 wb.close() 后删除。
        """
        try:
            return openpyxl.load_workbook(source_file, **kwargs), None
        except Exception:
            temp_path = self._make_repaired_copy(source_file)
            try:
                return openpyxl.load_workbook(temp_path, **kwargs), temp_path
            except Exception:
                if os.path.exists(temp_path):
                    os.remove(temp_path)
                raise
    
    def _process_file_group(self, save_file, rows):
        if not rows:
            return False
        source_file = rows[0].get('source_file', '')
        if not source_file or not os.path.exists(source_file):
            log_error(f'Source file not found: {source_file}')
            return False
        # B3-2：解析前限制源文件大小（防资源耗尽）
        if not check_source_file_size(source_file):
            log_error(f'Source file too large (>200 MB): {source_file}')
            return False
        # 文件类型与 Excel Tab 不符（如把 Word/PPT 文件填到 Excel Tab）时给出明确提示
        mismatch = document_kind_mismatch_message(source_file, 'excel')
        if mismatch:
            log_error(f'[ExcelWorker] Document kind mismatch: {source_file}')
            raise ValueError(mismatch)
        # 检测VBA宏（仅对.xls和.xlsm文件，不阻止处理，仅记录警告）
        file_ext = os.path.splitext(source_file)[1].lower()
        if file_ext in ('.xls', '.xlsm'):
            self._check_macros(source_file)
        save_dir = os.path.dirname(save_file)
        if save_dir and not os.path.exists(save_dir):
            os.makedirs(save_dir, exist_ok=True)
        first_sheet_name = None
        for row in rows:
            for item in row.get('data', []):
                sheet = item.get('sheet', '')
                if sheet and sheet.strip():
                    first_sheet_name = sheet.strip()
                    break
            if first_sheet_name:
                break
        temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        temp_path = temp_file.name
        temp_file.close()
        shutil.copy2(source_file, temp_path)
        # 修复畸形 docProps XML（部分模板 core.xml 根元素后有多余内容，openpyxl 会解析失败）
        self._repair_docprops_xml(temp_path)
        wb = openpyxl.load_workbook(temp_path)
        # 把 Sheet 页编号/名称解析为真实工作表名，保证边框样式来源与恢复目标一致
        resolved_sheet_name = self._resolve_sheet_name(wb, first_sheet_name)
        cell_styles = self._get_merged_cell_borders_from_template(source_file, resolved_sheet_name)
        try:
            if self.fill_method == 'placeholder':
                self._process_placeholder_mode(wb, rows)
            else:
                self._process_title_mode(wb, rows)
            if os.path.abspath(save_file) == os.path.abspath(source_file):
                save_temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
                save_temp_path = save_temp.name
                save_temp.close()
                wb.save(save_temp_path)
                wb.close()
                shutil.move(save_temp_path, save_file)
            else:
                wb.save(save_file)
                wb.close()
            self._apply_merged_cell_borders(save_file, cell_styles, resolved_sheet_name)
            return True
        except Exception as e:
            log_error(f'Error processing file: {e}')
            return False
        finally:
            try:
                if os.path.exists(temp_path):
                    os.remove(temp_path)
            except Exception:
                pass
    
    def _process_placeholder_mode(self, wb, rows):
        placeholder_values = defaultdict(list)
        placeholder_types = {}
        placeholder_sheets = {}
        for row in rows:
            is_inherited = row.get('is_inherited', False)
            for item in row.get('data', []):
                placeholder = item.get('placeholder', '')
                value = item.get('value', '')
                value_count = item.get('value_count', '')
                sheet = item.get('sheet', '')
                if placeholder:
                    is_multi_value = 'multi' in value_count.lower() or '多值' in value_count
                    if is_inherited and not is_multi_value:
                        continue
                    placeholder_values[placeholder].append(value)
                    if placeholder not in placeholder_types:
                        placeholder_types[placeholder] = value_count
                    if placeholder not in placeholder_sheets:
                        placeholder_sheets[placeholder] = sheet.strip() if sheet else ''
        
        placeholder_indices = defaultdict(int)
        
        worksheets_to_process = list(wb.worksheets)
        
        for ws in worksheets_to_process:
            for row_idx in range(1, ws.max_row + 1):
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if cell.data_type == 'f':
                        continue
                    if cell.value and isinstance(cell.value, str):
                        cell_value = cell.value
                        for placeholder, values in placeholder_values.items():
                            if placeholder in cell_value:
                                sheet_spec = placeholder_sheets.get(placeholder, '')
                                if sheet_spec:
                                    ws_list = self._get_worksheets_to_process(wb, sheet_spec)
                                    if ws not in ws_list:
                                        # 规格4.4：非法/超范围 Sheet 页跳过该列，避免误改活动工作表
                                        if not ws_list:
                                            log_warning(f'[ExcelWorker] Invalid sheet "{sheet_spec}", skipped')
                                            self._invalid_targets.append(str(sheet_spec))
                                        continue
                                value_count = placeholder_types.get(placeholder, '')
                                if 'multi' in value_count.lower() or '多值' in value_count:
                                    # 规格4.2：多值模式下同一单元格内的多个占位符依次消费各行值
                                    while placeholder in cell_value:
                                        idx = placeholder_indices[placeholder]
                                        if idx >= len(values):
                                            break
                                        new_value = values[idx]
                                        placeholder_indices[placeholder] = idx + 1
                                        cell_value = cell_value.replace(placeholder, str(new_value), 1)
                                elif 'image' in value_count.lower() or '图片' in value_count:
                                    if values:
                                        img_path = values[0]
                                        if os.path.exists(img_path):
                                            self._insert_image(ws, row_idx, col_idx, img_path)
                                            cell_value = cell_value.replace(placeholder, '')
                                else:
                                    if values:
                                        new_value = values[0]
                                        cell_value = cell_value.replace(placeholder, str(new_value))
                        cell.value = cell_value
    
    def _process_title_mode(self, wb, rows):
        for row_data in rows:
            is_inherited = row_data.get('is_inherited', False)
            multi_value_offset = row_data.get('multi_value_offset', 1)
            for item in row_data.get('data', []):
                title = item.get('placeholder', '')
                direction = item.get('direction', '')
                value_count = item.get('value_count', '')
                value = item.get('value', '')
                sheet = item.get('sheet', '')
                sheet = sheet.strip() if sheet else ''
                if not title:
                    continue
                is_multi_value = '多值' in value_count or 'multi' in value_count.lower()
                if is_inherited and not is_multi_value:
                    continue
                worksheets_to_process = self._get_worksheets_to_process(wb, sheet)
                coordinate_pos = None
                if self._is_direction_coordinate(direction):
                    parsed = self._parse_cell_address_with_sheet(title)
                    if parsed:
                        address_sheet_name, coordinate_pos = parsed
                        # 规格4.4：地址中的页名优先级高于 Sheet 页表头行
                        if address_sheet_name and address_sheet_name in wb.sheetnames:
                            worksheets_to_process = [wb[address_sheet_name]]
                    else:
                        coordinate_pos = self._parse_cell_address(title)
                if not worksheets_to_process:
                    # 规格4.4：非法/超范围 Sheet 页跳过该列，避免误改活动工作表
                    log_warning(f'[ExcelWorker] Invalid sheet "{sheet}", skipped')
                    self._invalid_targets.append(str(sheet))
                    continue
                for ws in worksheets_to_process:
                    if self._is_direction_coordinate(direction):
                        cell_positions = [coordinate_pos] if coordinate_pos else []
                    else:
                        cell_positions = self._find_all_title_cells(ws, title)
                    if not cell_positions:
                        continue
                    for cell_pos in cell_positions:
                        if is_multi_value:
                            target_row, target_col = self._get_target_position_with_offset(ws, cell_pos, direction, multi_value_offset)
                        else:
                            target_row, target_col = self._get_target_position(ws, cell_pos, direction)
                        # 规格4.7：目标单元格若位于合并区域内，取合并区域首个单元格
                        actual_row, actual_col = self._get_merge_start_cell(ws, target_row, target_col)
                        target_cell = ws.cell(row=actual_row, column=actual_col)
                        if target_cell.data_type == 'f':
                            continue
                        if 'image' in value_count.lower() or '图片' in value_count:
                            if value and os.path.exists(value):
                                self._insert_image(ws, actual_row, actual_col, value)
                        else:
                            if is_multi_value:
                                if value is not None and value != '':
                                    ws.cell(row=actual_row, column=actual_col, value=value)
                            else:
                                ws.cell(row=actual_row, column=actual_col, value=value)
    
    def _get_worksheets_to_process(self, wb, target_page):
        if not target_page:
            return wb.worksheets
        if target_page.isdigit():
            page_num = int(target_page)
            if 1 <= page_num <= len(wb.worksheets):
                return [wb.worksheets[page_num - 1]]
            # 规格4.4：非法/超范围 Sheet 页编号跳过该列，不再回退到活动工作表
            return []
        else:
            sheet_name = target_page
            if sheet_name in wb.sheetnames:
                return [wb[sheet_name]]
            return []
    
    def _get_merged_cell_borders_from_template(self, template_path, sheet_name=None):
        ns = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        rel_ns = {'rel': 'http://schemas.openxmlformats.org/package/2006/relationships'}
        try:
            with zipfile.ZipFile(template_path, 'r') as zf:
                sheet_target = 'xl/worksheets/sheet1.xml'
                if sheet_name:
                    try:
                        workbook_xml = zf.read('xl/workbook.xml').decode('utf-8')
                        wb_tree = ET.fromstring(workbook_xml)
                        sheets = wb_tree.findall('.//main:sheet', ns)
                        sheet_rids = {}
                        for s in sheets:
                            s_name = s.get('name', '')
                            s_rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id', '')
                            sheet_rids[s_name] = s_rid
                        rid = sheet_rids.get(sheet_name)
                        if rid:
                            rels_xml = zf.read('xl/_rels/workbook.xml.rels').decode('utf-8')
                            rels_tree = ET.fromstring(rels_xml)
                            for rel in rels_tree.findall('rel:Relationship', rel_ns):
                                if rel.get('Id') == rid:
                                    target = rel.get('Target', '')
                                    if target and not target.startswith('xl/'):
                                        target = 'xl/' + target
                                    sheet_target = target
                                    break
                    except Exception as e:
                        log_warning(f'Failed to resolve sheet name to file, using sheet1: {e}')
                sheet_xml = zf.read(sheet_target).decode('utf-8')
                styles_xml = zf.read('xl/styles.xml').decode('utf-8')
        except Exception as e:
            log_warning(f'Failed to read template XML: {e}')
            return {}
        try:
            styles_tree = ET.fromstring(styles_xml)
            borders_elem = styles_tree.find('.//main:borders', ns)
            cell_xfs_elem = styles_tree.find('.//main:cellXfs', ns)
        except Exception as e:
            log_warning(f'Failed to parse styles XML: {e}')
            return {}
        borders = borders_elem.findall('main:border', ns) if borders_elem is not None else []
        cell_xfs = cell_xfs_elem.findall('main:xf', ns) if cell_xfs_elem is not None else []
        def get_border_style(style_id):
            if style_id >= len(cell_xfs):
                return None
            xf = cell_xfs[style_id]
            border_id = int(xf.get('borderId', 0))
            if border_id >= len(borders):
                return None
            border = borders[border_id]
            def get_side(side_name):
                side = border.find(f'main:{side_name}', ns)
                if side is not None:
                    style = side.get('style')
                    if style:
                        return Side(style=style)
                return Side(style=None)
            return Border(
                left=get_side('left'),
                right=get_side('right'),
                top=get_side('top'),
                bottom=get_side('bottom')
            )
        cell_styles = {}
        for match in re.finditer(r'<c r="([A-Z]+\d+)" s="(\d+)"', sheet_xml):
            cell_ref = match.group(1)
            style_id = int(match.group(2))
            border = get_border_style(style_id)
            if border:
                cell_styles[cell_ref] = border
        return cell_styles
    
    def _resolve_sheet_name(self, wb, target_page):
        """把 Sheet 页编号或名称解析为真实工作表名；非法/为空返回 None（表示当前活动页/全部）"""
        if not target_page or not target_page.strip():
            return None
        target_page = target_page.strip()
        if target_page.isdigit():
            page_num = int(target_page)
            if 1 <= page_num <= len(wb.worksheets):
                return wb.worksheets[page_num - 1].title
            return None
        if target_page in wb.sheetnames:
            return target_page
        return None
    
    def _apply_merged_cell_borders(self, file_path, cell_styles, sheet_name=None):
        if not cell_styles:
            return
        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            for merged_range in ws.merged_cells.ranges:
                min_row, min_col = merged_range.min_row, merged_range.min_col
                max_row, max_col = merged_range.max_row, merged_range.max_col
                for row in range(min_row, max_row + 1):
                    for col in range(min_col, max_col + 1):
                        cell_ref = get_column_letter(col) + str(row)
                        if cell_ref in cell_styles:
                            cell = ws.cell(row=row, column=col)
                            cell.border = cell_styles[cell_ref]
            wb.save(file_path)
            wb.close()
            log_info(f'Applied merged cell borders to: {file_path}')
        except Exception as e:
            log_error(f'Failed to apply merged cell borders: {e}')
    
    def _get_merge_start_cell(self, ws, row, col):
        for merged_range in ws.merged_cells.ranges:
            if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
                return merged_range.min_row, merged_range.min_col
        return row, col
    
    def _get_merge_range_for_cell(self, ws, row, col):
        for merged_range in ws.merged_cells.ranges:
            if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
                return merged_range
        return None
    
    def _get_cell_bg_color(self, cell):
        """取单元格背景色的近似标识（用于“同背景色+同外框”区域判定）

        局限：主题色（type='theme'）使用硬编码的近似色表，未读取主题 tint/实际主题调色板，
        因此不同 tint 的同主题色会被视为相同颜色。该结果仅用于区域分组的启发式比较，
        不用于渲染，不影响单元格数据的写入。
        """
        if cell.fill and cell.fill.fill_type and cell.fill.fill_type != 'none':
            fg_color = cell.fill.fgColor
            if fg_color:
                if fg_color.type == 'rgb' and fg_color.rgb:
                    rgb_val = fg_color.rgb
                    if isinstance(rgb_val, str) and rgb_val != '00000000':
                        return rgb_val
                elif fg_color.type == 'theme' and fg_color.theme is not None:
                    theme_idx = fg_color.theme
                    theme_colors = [
                        'FFFFFF', '000000', 'E7E6E6', '44546A',
                        '5B9BD5', 'ED7D31', 'A5A5A5', 'FFC000',
                        '4472C4', '70AD47'
                    ]
                    if 0 <= theme_idx < len(theme_colors):
                        return theme_colors[theme_idx]
        return None
    
    def _has_real_border(self, border_style):
        if not border_style:
            return False
        return border_style not in ('none', 'default', None)
    
    def _get_cell_border_key(self, cell):
        borders = []
        if cell.border:
            if self._has_real_border(cell.border.left and cell.border.left.style):
                borders.append('L')
            if self._has_real_border(cell.border.right and cell.border.right.style):
                borders.append('R')
            if self._has_real_border(cell.border.top and cell.border.top.style):
                borders.append('T')
            if self._has_real_border(cell.border.bottom and cell.border.bottom.style):
                borders.append('B')
        return ''.join(sorted(borders))
    
    def _find_style_region(self, ws, start_row, start_col):
        """识别“相同背景色 + 共同非默认外框”的矩形区域（规格4.5）

        算法：以标题单元格为起点，沿其所在行/列向四个方向扩展，遇到不同背景色或
        真实边框则停止；随后要求区域边界上至少存在一条非默认外框线，否则返回 None。

        已知局限（启发式，非精确）：
        - 内部分隔边线只检查起始行/列方向上的相邻边线，区域内部其它位置的分隔线可能漏检，
          个别情况下会导致区域被判定得比实际偏大；
        - 主题色使用 `_get_cell_bg_color` 的近似色表，见其说明。
        若实际模板出现误判，应在扩展后补充“内部边线全扫描”与真实主题色读取。
        """
        start_cell = ws.cell(row=start_row, column=start_col)
        start_bg = self._get_cell_bg_color(start_cell)
        if not start_bg:
            return None
        min_row = start_row
        max_row = start_row
        min_col = start_col
        max_col = start_col
        for r in range(start_row - 1, 0, -1):
            current_cell = ws.cell(row=r + 1, column=start_col)
            if self._has_real_border(current_cell.border and current_cell.border.top and current_cell.border.top.style):
                break
            cell = ws.cell(row=r, column=start_col)
            if self._get_cell_bg_color(cell) == start_bg:
                min_row = r
            else:
                break
        for r in range(start_row + 1, ws.max_row + 1):
            current_cell = ws.cell(row=r - 1, column=start_col)
            if self._has_real_border(current_cell.border and current_cell.border.bottom and current_cell.border.bottom.style):
                break
            cell = ws.cell(row=r, column=start_col)
            if self._get_cell_bg_color(cell) == start_bg:
                max_row = r
            else:
                break
        for c in range(start_col - 1, 0, -1):
            current_cell = ws.cell(row=start_row, column=c + 1)
            if self._has_real_border(current_cell.border and current_cell.border.left and current_cell.border.left.style):
                break
            cell = ws.cell(row=start_row, column=c)
            if self._get_cell_bg_color(cell) == start_bg:
                min_col = c
            else:
                break
        for c in range(start_col + 1, ws.max_column + 1):
            current_cell = ws.cell(row=start_row, column=c - 1)
            if self._has_real_border(current_cell.border and current_cell.border.right and current_cell.border.right.style):
                break
            cell = ws.cell(row=start_row, column=c)
            if self._get_cell_bg_color(cell) == start_bg:
                max_col = c
            else:
                break
        has_outer_border = False
        for c in range(min_col, max_col + 1):
            top_cell = ws.cell(row=min_row, column=c)
            if self._has_real_border(top_cell.border and top_cell.border.top and top_cell.border.top.style):
                has_outer_border = True
                break
            bottom_cell = ws.cell(row=max_row, column=c)
            if self._has_real_border(bottom_cell.border and bottom_cell.border.bottom and bottom_cell.border.bottom.style):
                has_outer_border = True
                break
        if not has_outer_border:
            for r in range(min_row, max_row + 1):
                left_cell = ws.cell(row=r, column=min_col)
                if self._has_real_border(left_cell.border and left_cell.border.left and left_cell.border.left.style):
                    has_outer_border = True
                    break
                right_cell = ws.cell(row=r, column=max_col)
                if self._has_real_border(right_cell.border and right_cell.border.right and right_cell.border.right.style):
                    has_outer_border = True
                    break
        if not has_outer_border:
            return None
        return (min_row, max_row, min_col, max_col)
    
    def _is_direction_down(self, direction):
        if not direction:
            return False
        direction_lower = direction.lower()
        return 'down' in direction_lower or '向下' in direction or translate('direction_down') in direction
    
    def _is_direction_right(self, direction):
        if not direction:
            return False
        direction_lower = direction.lower()
        return 'right' in direction_lower or '向右' in direction or translate('direction_right') in direction
    
    def _is_direction_coordinate(self, direction):
        if not direction:
            return False
        direction_lower = direction.lower()
        return 'coordinate' in direction_lower or '坐标' in direction or translate('direction_coordinate') in direction
    
    def _get_target_position_with_offset(self, ws, cell_pos, direction, offset):
        row, col = cell_pos
        if self._is_direction_coordinate(direction):
            return (row, col)
        merge_range = self._get_merge_range_for_cell(ws, row, col)
        if merge_range:
            if self._is_direction_down(direction):
                return (merge_range.max_row + offset, merge_range.min_col)
            elif self._is_direction_right(direction):
                return (merge_range.min_row, merge_range.max_col + offset)
            else:
                return (merge_range.max_row + offset, merge_range.min_col)
        style_region = self._find_style_region(ws, row, col)
        if style_region:
            min_row, max_row, min_col, max_col = style_region
            if self._is_direction_down(direction):
                return (max_row + offset, min_col)
            elif self._is_direction_right(direction):
                return (min_row, max_col + offset)
            else:
                return (max_row + offset, min_col)
        if self._is_direction_down(direction):
            return (row + offset, col)
        elif self._is_direction_right(direction):
            return (row, col + offset)
        else:
            return (row + offset, col)
    
    def _get_target_position(self, ws, cell_pos, direction):
        row, col = cell_pos
        if self._is_direction_coordinate(direction):
            return (row, col)
        merge_range = self._get_merge_range_for_cell(ws, row, col)
        if merge_range:
            if self._is_direction_down(direction):
                return (merge_range.max_row + 1, merge_range.min_col)
            elif self._is_direction_right(direction):
                return (merge_range.min_row, merge_range.max_col + 1)
            else:
                return (merge_range.max_row + 1, merge_range.min_col)
        style_region = self._find_style_region(ws, row, col)
        if style_region:
            min_row, max_row, min_col, max_col = style_region
            if self._is_direction_down(direction):
                return (max_row + 1, min_col)
            elif self._is_direction_right(direction):
                return (min_row, max_col + 1)
            else:
                return (max_row + 1, min_col)
        if self._is_direction_down(direction):
            return (row + 1, col)
        elif self._is_direction_right(direction):
            return (row, col + 1)
        else:
            return (row + 1, col)
    
    def _clean_title(self, title):
        title = title.replace('\r\n', '\n')
        title = title.strip()
        return title
    
    def _parse_cell_address(self, address):
        address = address.strip().upper()
        match = re.match(r'^([A-Z]+)(\d+)$', address)
        if not match:
            return None
        col_str = match.group(1)
        row_num = int(match.group(2))
        col_num = 0
        for char in col_str:
            col_num = col_num * 26 + (ord(char) - ord('A') + 1)
        return (row_num, col_num)
    
    def _parse_cell_address_with_sheet(self, address):
        address = address.strip()
        if '!' in address:
            parts = address.split('!', 1)
            sheet_name = parts[0].strip().strip("'")
            cell_address = parts[1].strip()
            cell_pos = self._parse_cell_address(cell_address)
            if cell_pos:
                return (sheet_name, cell_pos)
        return None
    
    def _get_worksheet_by_target_page(self, wb, target_page):
        if not target_page or not target_page.strip():
            return None
        target_page = target_page.strip()
        if target_page.isdigit():
            page_num = int(target_page)
            if 1 <= page_num <= len(wb.worksheets):
                return wb.worksheets[page_num - 1]
        else:
            sheet_name = target_page
            if sheet_name in wb.sheetnames:
                return wb[sheet_name]
        return None
    
    def _find_title_cell(self, ws, title):
        title_clean = self._clean_title(title)
        title_variants = [t.strip() for t in title_clean.split('\n') if t.strip()]
        title_variants.append(title_clean.replace('\n', ' '))
        for row_idx in range(1, ws.max_row + 1):
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value:
                    cell_value = str(cell.value).strip()
                    cell_value_normalized = re.sub(r'\s+', ' ', cell_value)
                    for variant in title_variants:
                        variant_normalized = re.sub(r'\s+', ' ', variant)
                        if cell_value_normalized == variant_normalized:
                            return (row_idx, col_idx)
        return None
    
    def _find_all_title_cells(self, ws, title):
        title_clean = self._clean_title(title)
        title_variants = [t.strip() for t in title_clean.split('\n') if t.strip()]
        title_variants.append(title_clean.replace('\n', ' '))
        positions = []
        for row_idx in range(1, ws.max_row + 1):
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value:
                    cell_value = str(cell.value).strip()
                    cell_value_normalized = re.sub(r'\s+', ' ', cell_value)
                    for variant in title_variants:
                        variant_normalized = re.sub(r'\s+', ' ', variant)
                        if cell_value_normalized == variant_normalized:
                            positions.append((row_idx, col_idx))
                            break
        return positions
    
    def _insert_image(self, ws, row, col, img_path):
        if not HAS_PIL:
            log_error('PIL library not installed for image processing')
            return
        # B3-2：解析前限制图片源文件大小（防资源耗尽）
        if not check_source_file_size(img_path):
            log_error(f'Image file too large (>200 MB): {img_path}')
            return
        try:
            img_params = self.params.get('image_params', {})
            max_width = img_params.get('width', 200)
            max_height = img_params.get('height', 200)
            pil_img = Image.open(img_path)
            orig_width, orig_height = pil_img.size
            ratio = min(max_width / orig_width, max_height / orig_height)
            if ratio < 1:
                new_width = int(orig_width * ratio)
                new_height = int(orig_height * ratio)
            else:
                new_width = orig_width
                new_height = orig_height
            pil_img_resized = pil_img.resize((new_width, new_height), Image.Resampling.LANCZOS)
            from io import BytesIO
            img_stream = BytesIO()
            pil_img_resized.save(img_stream, format='PNG')
            img_stream.seek(0)
            xl_img = XLImage(img_stream)
            xl_img.width = new_width
            xl_img.height = new_height
            cell_addr = f'{get_column_letter(col)}{row}'
            xl_img.anchor = cell_addr
            ws.add_image(xl_img)
            pil_img.close()
        except Exception as e:
            import traceback
            log_error(f'Failed to insert image: {e}')
            log_error(traceback.format_exc())
    
    def _execute_get_file_info(self):
        if not self.data:
            self._show_error(translate('msg_no_data'))
            return False
        
        grid = None
        if self.tool_instance and hasattr(self.tool_instance, 'grids'):
            grid = self.tool_instance.grids.get('excel')
        
        success_count = 0
        error_count = 0
        
        for idx, row_data in enumerate(self.data):
            source_file = row_data.get('source_file', '')
            if not source_file or not source_file.strip():
                continue
            
            if not os.path.exists(source_file):
                log_warning(f'File not found: {source_file}')
                self.log_info_row_failure(row_data, 'file not found')
                error_count += 1
                continue
            # 文件类型与 Excel Tab 不符时给出明确提示
            mismatch = document_kind_mismatch_message(source_file, 'excel')
            if mismatch:
                log_error(f'[ExcelWorker] Document kind mismatch, skip get info: {source_file}')
                self.log_info_row_failure(row_data, mismatch)
                error_count += 1
                continue
            try:
                file_size = os.path.getsize(source_file)
                if file_size < 1024:
                    file_size_str = f'{file_size} B'
                elif file_size < 1024 * 1024:
                    file_size_str = f'{file_size / 1024:.1f} KB'
                else:
                    file_size_str = f'{file_size / (1024 * 1024):.1f} MB'
                
                file_ext = os.path.splitext(source_file)[1].lower()
                
                if file_ext == '.xls':
                    info_values = self._get_xls_file_info(source_file, file_size_str)
                else:
                    info_values = self._get_xlsx_file_info(source_file, file_size_str)
                
                if info_values and grid:
                    row_index = row_data.get('row_index', 0)
                    for col_idx, value in enumerate(info_values):
                        grid.SetCellValue(row_index, grid.COL_DATA_START + col_idx, value)
                
                if info_values:
                    log_info(f'Excel info extracted: {source_file}')
                    success_count += 1
                else:
                    self.log_info_row_failure(row_data, 'failed to extract info')
                    error_count += 1
                
            except Exception as e:
                log_error(f'Error extracting Excel info from {source_file}: {e}')
                self.log_info_row_failure(row_data, str(e))
                error_count += 1
        
        log_info(f'Get Excel Info completed: {success_count} success, {error_count} errors')
        
        # 获取文档信息：仅在存在未成功的行时弹窗提示，全部成功时静默完成
        if error_count > 0:
            message = translate('msg_info_failed_rows', count=error_count)
            wx.MessageBox(message, translate('msg_warning'), wx.OK | wx.ICON_WARNING)
        
        return True
    
    def _get_xlsx_file_info(self, source_file, file_size_str):
        if not HAS_OPENPYXL:
            log_error('openpyxl library not installed')
            return None
        # B3-2：解析前限制源文件大小（防资源耗尽）
        if not check_source_file_size(source_file):
            log_error(f'Source file too large (>200 MB): {source_file}')
            return None
        
        wb, temp_path = self._load_workbook_repairing(source_file, read_only=True)
        sheet_count = len(wb.worksheets)
        
        ws = wb.active
        max_row = ws.max_row
        max_col = ws.max_column
        
        props = wb.properties
        title = props.title or ''
        author = props.creator or translate('pdf_info_not_available')
        created = props.created
        modified = props.modified
        last_modified_by = props.lastModifiedBy or translate('pdf_info_not_available')
        
        if not title:
            title = os.path.splitext(os.path.basename(source_file))[0]
        
        file_stat = os.stat(source_file)
        file_created = file_stat.st_ctime
        file_modified = file_stat.st_mtime
        
        created_str = translate('pdf_info_not_available')
        if created:
            try:
                created_str = created.strftime('%Y-%m-%d')
            except Exception:
                pass
        if created_str == translate('pdf_info_not_available'):
            try:
                from datetime import datetime
                created_str = datetime.fromtimestamp(file_created).strftime('%Y-%m-%d')
            except (ValueError, AttributeError):
                pass
        
        modified_str = translate('pdf_info_not_available')
        if modified:
            try:
                modified_str = modified.strftime('%Y-%m-%d')
            except Exception:
                pass
        if modified_str == translate('pdf_info_not_available'):
            try:
                from datetime import datetime
                modified_str = datetime.fromtimestamp(file_modified).strftime('%Y-%m-%d')
            except (ValueError, AttributeError):
                pass
        
        wb.close()
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass
        
        return [
            file_size_str,
            str(sheet_count),
            str(max_row),
            str(max_col),
            title,
            author,
            created_str,
            modified_str,
            last_modified_by
        ]
    
    def _get_xls_file_info(self, source_file, file_size_str):
        if not HAS_XLRD:
            log_error('xlrd library not installed, cannot read .xls files')
            return None
        # B3-2：解析前限制源文件大小（与 xlsx 分支保持一致）
        if not check_source_file_size(source_file):
            log_error(f'Source file too large (>200 MB): {source_file}')
            return None
        
        try:
            wb = xlrd.open_workbook(source_file)
            sheet_count = wb.nsheets
            
            ws = wb.sheet_by_index(0)
            max_row = ws.nrows
            max_col = ws.ncols
            
            title = ''
            author = translate('pdf_info_not_available')
            created_str = translate('pdf_info_not_available')
            modified_str = translate('pdf_info_not_available')
            last_modified_by = translate('pdf_info_not_available')
            
            try:
                props = wb.props
                if hasattr(props, 'title') and props.title:
                    title = props.title
                if hasattr(props, 'author') and props.author:
                    author = props.author
                if hasattr(props, 'created') and props.created:
                    try:
                        created_str = xlrd.xldate_as_datetime(props.created, wb.datemode).strftime('%Y-%m-%d')
                    except (ValueError, AttributeError):
                        pass
                if hasattr(props, 'modified') and props.modified:
                    try:
                        modified_str = xlrd.xldate_as_datetime(props.modified, wb.datemode).strftime('%Y-%m-%d')
                    except (ValueError, AttributeError):
                        pass
                if hasattr(props, 'last_modified_by') and props.last_modified_by:
                    last_modified_by = props.last_modified_by
            except Exception as e:
                log_warning(f'Failed to read .xls properties: {e}')
            
            if not title:
                title = os.path.splitext(os.path.basename(source_file))[0]
            
            file_stat = os.stat(source_file)
            file_created = file_stat.st_ctime
            file_modified = file_stat.st_mtime
            
            if created_str == translate('pdf_info_not_available'):
                try:
                    from datetime import datetime
                    created_str = datetime.fromtimestamp(file_created).strftime('%Y-%m-%d')
                except (ValueError, AttributeError):
                    pass
            
            if modified_str == translate('pdf_info_not_available'):
                try:
                    from datetime import datetime
                    modified_str = datetime.fromtimestamp(file_modified).strftime('%Y-%m-%d')
                except (ValueError, AttributeError):
                    pass
            
            return [
                file_size_str,
                str(sheet_count),
                str(max_row),
                str(max_col),
                title,
                author,
                created_str,
                modified_str,
                last_modified_by
            ]
        except Exception as e:
            log_error(f'Error reading .xls file {source_file}: {e}')
            return None
    
    def _show_results(self):
        success_count = sum(1 for r in self.results if r.get('success'))
        total_count = len(self.results)
        if self.tool_instance and hasattr(self.tool_instance, 'grids'):
            grid = self.tool_instance.grids.get('excel')
            if grid:
                for r in self.results:
                    if r.get('warning'):
                        grid.set_result(r['file'], r['warning'])
                    elif r.get('success'):
                        grid.set_result(r['file'], translate('result_success'))
                    else:
                        error_msg = r.get('error', '')
                        short_error = self._get_short_error(error_msg)
                        grid.set_result(r['file'], f'{translate("result_failed")}: {short_error}')
        message = self._build_completion_message(success_count, total_count)
        def _do_show():
            wx.MessageBox(message, translate('tool_name'), wx.OK | wx.ICON_INFORMATION)
        wx.CallAfter(_do_show)
    
    def _get_short_error(self, error_msg):
        if not error_msg:
            return translate('result_error_unknown')
        # 文件类型与 Tab 不符：直接显示完整提示，避免被归类为泛化错误
        for key in ('msg_wrong_tab_use_word', 'msg_wrong_tab_use_excel', 'msg_wrong_tab_use_ppt'):
            if error_msg == translate(key):
                return error_msg
        error_lower = error_msg.lower()
        if 'permission' in error_lower or '权限' in error_msg:
            return translate('result_error_permission')
        if 'not found' in error_lower or '不存在' in error_msg:
            return translate('result_error_not_found')
        if 'invalid' in error_lower or '无效' in error_msg:
            return translate('result_error_invalid')
        if 'timeout' in error_lower or '超时' in error_msg:
            return translate('result_error_timeout')
        if 'memory' in error_lower or '内存' in error_msg:
            return translate('result_error_memory')
        return translate('result_error_other')
    
    def _resolve_preview_sheet(self, wb, rows_data):
        """确定预览展示的工作表：优先 Sheet 表头；坐标模式支持地址页名；否则活动页（规格4.4）"""
        for row in rows_data:
            for item in row.get('data', []):
                spec = (item.get('sheet') or '').strip()
                direction = item.get('direction', '')
                if self._is_direction_coordinate(direction):
                    parsed = self._parse_cell_address_with_sheet(item.get('placeholder', ''))
                    if parsed and parsed[0] in wb.sheetnames:
                        spec = parsed[0]
                if spec:
                    sheets = self._get_worksheets_to_process(wb, spec)
                    if sheets:
                        return sheets[0]
        return wb.active
    
    def preview(self, rows_data):
        if not HAS_OPENPYXL:
            return None
        if not rows_data:
            return None
        if isinstance(rows_data, dict):
            rows_data = [rows_data]
        source_file = rows_data[0].get('source_file', '')
        save_file = rows_data[0].get('save_file', '')
        if not source_file or not os.path.exists(source_file):
            self.log_group_result(save_file, rows_data, False, error='Source file not found')
            return None
        # B3-2：解析前限制源文件大小（防资源耗尽）
        if not check_source_file_size(source_file):
            log_error(f'Source file too large (>200 MB), skip preview: {source_file}')
            self.log_group_result(save_file, rows_data, False, error='Source file too large (>200 MB)')
            return None
        # 文件类型与 Excel Tab 不符时给出明确提示（规格：文件类型与 Tab 一一对应）
        mismatch = document_kind_mismatch_message(source_file, 'excel')
        if mismatch:
            log_error(f'[ExcelWorker] Document kind mismatch, skip preview: {source_file}')
            self.log_group_result(save_file, rows_data, False, error=mismatch)
            return {'type': 'excel', 'error': mismatch}
        temp_paths = []
        try:
            wb, temp_path = self._load_workbook_repairing(source_file, read_only=False)
            if temp_path:
                temp_paths.append(temp_path)
            ws = self._resolve_preview_sheet(wb, rows_data)
            if self.fill_method == 'placeholder':
                self._process_placeholder_mode(wb, rows_data)
            else:
                self._process_title_mode(wb, rows_data)
            max_preview_rows = min(ws.max_row + 1, 100)
            max_preview_cols = min(ws.max_column + 1, 50)
            format_info = self._extract_format_info(ws, max_preview_rows, max_preview_cols)
            max_row = ws.max_row
            max_col = ws.max_column
            wb.close()
            wb_data, temp_path_data = self._load_workbook_repairing(source_file, data_only=True)
            if temp_path_data:
                temp_paths.append(temp_path_data)
            ws_data = self._resolve_preview_sheet(wb_data, rows_data)
            if self.fill_method == 'placeholder':
                self._process_placeholder_mode(wb_data, rows_data)
            else:
                self._process_title_mode(wb_data, rows_data)
            preview_data = []
            for row_idx in range(1, max_preview_rows):
                cell_row = []
                for col_idx in range(1, max_preview_cols):
                    cell = ws_data.cell(row=row_idx, column=col_idx)
                    value = cell.value
                    if value is None:
                        value = ''
                    cell_row.append(str(value))
                preview_data.append(cell_row)
            wb_data.close()
            self.log_group_result(save_file, rows_data, True)
            return {
                'type': 'excel',
                'data': preview_data,
                'max_row': max_row,
                'max_col': max_col,
                'format': format_info
            }
        except Exception as e:
            import traceback
            log_error(traceback.format_exc())
            self.log_group_result(save_file, rows_data, False, error=str(e))
            return None
        finally:
            for temp_path in temp_paths:
                try:
                    if os.path.exists(temp_path):
                        os.remove(temp_path)
                except Exception:
                    pass
    
    def _extract_format_info(self, ws, max_rows, max_cols):
        format_info = {
            'row_heights': {},
            'col_widths': {},
            'hidden_rows': [],
            'hidden_cols': [],
            'merged_cells': [],
            'cell_backgrounds': {},
            'cell_borders': {},
            'cell_alignments': {},
            'cell_font_colors': {}
        }
        try:
            for row_idx in range(1, max_rows + 1):
                try:
                    row_dim = ws.row_dimensions[row_idx]
                    if row_dim.height:
                        format_info['row_heights'][row_idx] = row_dim.height
                    if row_dim.hidden:
                        format_info['hidden_rows'].append(row_idx)
                except Exception:
                    pass
            for col_idx in range(1, max_cols + 1):
                try:
                    col_letter = get_column_letter(col_idx)
                    col_width = None
                    is_hidden = False
                    if col_letter in ws.column_dimensions:
                        col_dim = ws.column_dimensions[col_letter]
                        if col_dim.width is not None:
                            col_width = col_dim.width
                        if col_dim.hidden:
                            is_hidden = True
                    if col_width is None:
                        col_width = 8.43
                    format_info['col_widths'][col_idx] = col_width
                    if is_hidden:
                        format_info['hidden_cols'].append(col_idx)
                except Exception:
                    format_info['col_widths'][col_idx] = 8.43
            try:
                for merged_range in ws.merged_cells.ranges:
                    format_info['merged_cells'].append({
                        'min_row': merged_range.min_row,
                        'max_row': merged_range.max_row,
                        'min_col': merged_range.min_col,
                        'max_col': merged_range.max_col
                    })
            except Exception:
                pass
            for row_idx in range(1, max_rows + 1):
                for col_idx in range(1, max_cols + 1):
                    try:
                        cell = ws.cell(row=row_idx, column=col_idx)
                        if cell.fill and cell.fill.fill_type and cell.fill.fill_type != 'none':
                            fg_color = cell.fill.fgColor
                            bg_color = None
                            if fg_color:
                                if fg_color.type == 'rgb' and fg_color.rgb:
                                    rgb_val = fg_color.rgb
                                    if isinstance(rgb_val, str) and rgb_val != '00000000' and len(rgb_val) >= 6:
                                        bg_color = rgb_val
                                elif fg_color.type == 'theme' and fg_color.theme is not None:
                                    try:
                                        theme_idx = fg_color.theme
                                        tint = fg_color.tint if fg_color.tint else 0.0
                                        theme_colors = [
                                            'FFFFFF', '000000', 'E7E6E6', '44546A',
                                            '5B9BD5', 'ED7D31', 'A5A5A5', 'FFC000',
                                            '4472C4', '70AD47'
                                        ]
                                        if 0 <= theme_idx < len(theme_colors):
                                            bg_color = theme_colors[theme_idx]
                                        else:
                                            bg_color = 'FFFFFF'
                                    except (ValueError, AttributeError):
                                        bg_color = 'FFFFFF'
                            if bg_color:
                                format_info['cell_backgrounds'][f'{row_idx},{col_idx}'] = bg_color
                        if cell.border:
                            has_border = (
                                cell.border.left.style or
                                cell.border.right.style or
                                cell.border.top.style or
                                cell.border.bottom.style
                            )
                            if has_border:
                                border_info = {}
                                if cell.border.left and cell.border.left.style:
                                    border_info['left'] = True
                                if cell.border.right and cell.border.right.style:
                                    border_info['right'] = True
                                if cell.border.top and cell.border.top.style:
                                    border_info['top'] = True
                                if cell.border.bottom and cell.border.bottom.style:
                                    border_info['bottom'] = True
                                if border_info:
                                    format_info['cell_borders'][f'{row_idx},{col_idx}'] = border_info
                        if cell.alignment:
                            align_info = {}
                            if cell.alignment.horizontal:
                                align_info['h'] = cell.alignment.horizontal
                            if cell.alignment.vertical:
                                align_info['v'] = cell.alignment.vertical
                            if align_info:
                                format_info['cell_alignments'][f'{row_idx},{col_idx}'] = align_info
                        if cell.font and cell.font.color:
                            font_color = cell.font.color
                            text_color = None
                            if font_color.type == 'rgb' and font_color.rgb:
                                rgb_val = font_color.rgb
                                if isinstance(rgb_val, str) and rgb_val != '00000000' and len(rgb_val) >= 6:
                                    text_color = rgb_val
                            elif font_color.type == 'theme' and font_color.theme is not None:
                                try:
                                    theme_idx = font_color.theme
                                    theme_colors = [
                                        'FFFFFF', '000000', 'E7E6E6', '44546A',
                                        '5B9BD5', 'ED7D31', 'A5A5A5', 'FFC000',
                                        '4472C4', '70AD47'
                                    ]
                                    if 0 <= theme_idx < len(theme_colors):
                                        text_color = theme_colors[theme_idx]
                                except (ValueError, AttributeError):
                                    pass
                            if text_color and text_color != '000000':
                                format_info['cell_font_colors'][f'{row_idx},{col_idx}'] = text_color
                    except (ValueError, AttributeError):
                        pass
        except Exception as e:
            log_error(f'Error extracting format info: {e}')
        return format_info
